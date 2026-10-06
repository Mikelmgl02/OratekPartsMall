"""Private pricing must never reach clients, competitors or Oratek staff.

Each slice that stores private pricing data seeds it in seed_private_markers(); the canary then calls every
client-reachable endpoint and fails if any marker leaks. The key-set pins fail on any new client-facing key.
"""
import io
import json
import logging
import uuid
from decimal import Decimal

from django.contrib import admin
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from rest_framework.test import APITestCase

from . import test_deals as deal_tests
from . import test_requests as request_tests
from .management import ManagedInventorySerializer
from .models import Membership, SupplierItem, User
from .price_import_models import PriceImportBatch, PriceImportJob
from .pricing_models import (ClientPricingProfile, PriceChange, PriceList, PriceListEntry, PricingAuditEvent, PricingRule, SupplierItemPricing,
                             SupplierPricingSettings, record_pricing_event)
from .pricing_services import set_prices
from .quote_draft_models import DealQuotationAudit, DealQuotationDraft, DealQuotationDraftLine, DealQuotationLineAudit
from .request_models import DealCommand, SupplierRequest
from .serializers import SupplierItemSerializer

SUMMARY_KEYS = {'id', 'reference', 'status', 'version', 'created_at', 'updated_at', 'handshaked_at', 'reviewed_at', 'line_count',
                'unit_count', 'quotation_revision', 'quoted_unit_count', 'quotation_total', 'quotation_currency'}
SENT_SUMMARY_KEYS = SUMMARY_KEYS | {'supplier', 'submission_id'}
DEAL_KEYS = {'client', 'supplier', 'quotations', 'quotation', 'events'}
QUOTATION_KEYS = {'id', 'revision', 'currency', 'terms', 'total', 'created_at', 'supplier_confirmed_at', 'client_confirmed_at', 'lines'}
QUOTATION_LINE_KEYS = {'order_line_id', 'codigo', 'sku', 'description', 'quantity', 'unit_price', 'total'}
EVENT_KEYS = {'id', 'kind', 'description', 'account_name', 'actor_name', 'created_at', 'quotation_id'}
CLIENT_LINE_KEYS = {'id', 'part_id', 'sku', 'name', 'codigo', 'brand', 'description', 'quantity'}
SUPPLIER_LINE_KEYS = CLIENT_LINE_KEYS | {'supplier_item_id', 'supplier_invent_id', 'stock'}
MARKERS = ['777.77', '666.66', '555.55', 'SECRETO-REGLA', 'SECRETO-PERFIL', 'SECRETO-BORRADOR', 'SECRETO-IA', 'SECRETO-AUDITORIA',
           'SECRETO-LINEA', 'LISTA-SECRETA', 'SECRETA_X', 'SECRETO-ERP', 'SECRETO-NOTA']
XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


class LogCapture(logging.Handler):
    """Every log record (INFO and above) emitted while the canary runs; SQL debug logging is not application logging."""
    def __init__(self):
        super().__init__(logging.INFO)
        self.lines = []

    def emit(self, record):
        self.lines.append(f'{record.name} {record.getMessage()} {record.args!r}')


class PricingPrivacyTests(APITestCase):
    setUp = request_tests.SupplierRequestWorkflowTests.setUp
    account = request_tests.SupplierRequestWorkflowTests.account
    item = request_tests.SupplierRequestWorkflowTests.item
    url = request_tests.SupplierRequestWorkflowTests.url
    payload = request_tests.SupplierRequestWorkflowTests.payload
    submit = request_tests.SupplierRequestWorkflowTests.submit
    order = deal_tests.DealWorkflowTests.order
    review = deal_tests.DealWorkflowTests.review
    action = deal_tests.DealWorkflowTests.action
    messages = deal_tests.DealWorkflowTests.messages
    trace_paths = ()
    import_job = uuid.UUID(int=0)

    def quote(self, order, price='13.00'):
        self.client.force_authenticate(self.seller_a)
        response = self.action(order, 'quote', lines=[{'order_line_id': str(line.pk), 'quantity': line.quantity, 'unit_price': price}
                                                      for line in order.lines.all()], terms='ENTREGA EN 48 HORAS')
        self.assertEqual(response.status_code, 200)
        return response

    def seed_private_markers(self, order):
        """Real private rows: price lists, entries, LINEA, floors and their history (S3); the pair profile (internal notes, ERP code) and
        commercial rules with a net price (S5); assistant output arrives with S7. Drafts only exist while the supplier prepares a revision, so
        seed_private_draft() adds one later."""
        SupplierPricingSettings.objects.create(supplier=self.supplier_a, default_currency='PAB', config_min_permission='manager',
                                               accept_shortfall_policy='allow', assistant_enabled=True, updated_by=self.seller_a)
        Membership.objects.filter(user=self.seller_a, account=self.supplier_a).update(permission='manager')
        default = PriceList.objects.create(supplier=self.supplier_a, code='SECRETA_X', name='LISTA-SECRETA', currency='USD', is_default=True, created_by=self.seller_a)
        set_prices(self.supplier_a, self.seller_a, changes=[{'list_id': default.pk, 'supplier_item_id': item.pk, 'unit_price': Decimal('777.77'), 'expected_revision': None}
                                                            for item in (self.item_a, self.item_a2)],
                   items=[{'supplier_item_id': self.item_a.pk, 'discount_group': 'SECRETO-LINEA', 'floor_price': Decimal('666.66'), 'expected_revision': None}],
                   reference='op:SECRETO-AUDITORIA')
        self.assertEqual((PriceListEntry.objects.count(), SupplierItemPricing.objects.count(), PriceChange.objects.count()), (2, 1, 4))
        record_pricing_event(self.supplier_a, self.seller_a, 'prices_edited', client=self.client_account, order=order, object_id='SECRETO-AUDITORIA',
                             payload={'list_price': '777.77', 'floor_price': '666.66', 'net_price': '555.55', 'rule': 'SECRETO-REGLA',
                                      'profile_note': 'SECRETO-PERFIL', 'draft_note': 'SECRETO-BORRADOR', 'assistant': 'SECRETO-IA'})
        self.seed_private_profile()
        stored = json.dumps([list(PricingAuditEvent.objects.values('object_id', 'payload')), list(PriceList.objects.values('code', 'name')),
                             list(PriceChange.objects.values('new_value', 'new_text')), list(SupplierItemPricing.objects.values('discount_group', 'floor_price')),
                             list(ClientPricingProfile.objects.values('internal_notes', 'customer_code')), list(PricingRule.objects.values('name', 'value', 'note'))],
                            default=str)
        self.assertTrue(all(marker in stored for marker in MARKERS), 'The canary must seed every marker it looks for.')
        self.seed_private_import()

    def seed_private_profile(self):
        """The pair profile and the supplier's rules (S5), saved through the supplier's own endpoints: an internal note, an ERP client code,
        a client net price and a general rule, all private."""
        self.client.force_authenticate(self.seller_a)
        base, client = f'/api/v1/accounts/{self.supplier_a.pk}', str(self.client_account.pk)
        profile = self.client.post(f'{base}/clients/{client}/profile/', {'operation_id': str(uuid.uuid4()), 'expected_version': 0, 'internal_notes': 'SECRETO-PERFIL',
                                                                         'customer_code': 'SECRETO-ERP', 'discount_percent': '3.00'}, format='json')
        self.assertEqual(profile.status_code, 200)
        for body in [{'name': 'SECRETO-REGLA', 'scope': 'client', 'client_id': client, 'target': 'item', 'item_id': str(self.item_a.pk), 'kind': 'net_price',
                      'value': '555.55', 'currency': 'USD', 'note': 'SECRETO-NOTA'},
                     {'name': 'SECRETO-REGLA GENERAL', 'scope': 'all', 'target': 'line', 'target_value': 'SECRETO-LINEA', 'kind': 'discount', 'value': '12.00'}]:
            self.assertEqual(self.client.post(f'{base}/pricing-rules/', {'operation_id': str(uuid.uuid4()), **body}, format='json').status_code, 201)
        self.rule = PricingRule.objects.get(name='SECRETO-REGLA')
        simulated = self.client.post(f'{base}/pricing/simulate/', {'client_id': client, 'items': [{'supplier_item_id': str(self.item_a.pk), 'quantity': 2}]}, format='json')
        self.assertEqual(simulated.status_code, 200)
        self.assertTrue(all(marker in simulated.content.decode() for marker in ['555.55', 'SECRETO-REGLA', '777.77']))

    def price_file(self, rows):
        book = Workbook()
        book.active.title = 'PRECIOS'
        for row in [['ID_INVENTARIO_PROVEEDOR', 'LINEA', 'PRECIO_MINIMO', 'PRECIO'], *rows]:
            book.active.append(row)
        output = io.BytesIO()
        book.save(output)
        return SimpleUploadedFile('PRECIOS.xlsx', output.getvalue())

    def seed_private_import(self):
        """A real price import preview (S4): its job keeps a price change and a rejected row with the private values; nothing is applied."""
        self.client.force_authenticate(self.seller_a)
        response = self.client.post(f'/api/v1/accounts/{self.supplier_a.pk}/prices/import/', {'file': self.price_file([
            ['A-001', 'SECRETO-LINEA', '666.66', '777.77'], ['A-002', '', '', '555.55'], ['NO-EXISTE', 'SECRETO-LINEA', '666.66', '777.77']])}, format='multipart')
        self.assertEqual((response.status_code, response.data['summary']['decreases'], response.data['summary']['rejected_rows']), (200, 1, 1))
        self.import_job = response.data['job_id']
        stored = json.dumps(list(PriceImportJob.objects.values('preview', 'rejected_rows')) + list(PriceImportBatch.objects.values('data')), default=str)
        self.assertTrue(all(marker in stored for marker in ['777.77', '666.66', '555.55', 'SECRETO-LINEA']))

    def seed_private_draft(self, order):
        """A real server draft (S1) with a private price and internal note, saved through the supplier's own endpoint."""
        self.client.force_authenticate(self.seller_a)
        response = self.client.post(f'/api/v1/accounts/{self.supplier_a.pk}/requests/{order.pk}/draft/', {
            'save_id': str(uuid.uuid4()), 'expected_draft_version': 0, 'terms': 'BORRADOR 666.66',
            'lines': [{'order_line_id': str(line.pk), 'unit_price': '777.77', 'note': 'SECRETO-BORRADOR'} for line in order.lines.all()]}, format='json')
        self.assertEqual(response.status_code, 200)
        self.assertIn('SECRETO-BORRADOR', response.content.decode())
        self.assertTrue(DealQuotationDraftLine.objects.filter(note='SECRETO-BORRADOR', unit_price='777.77').exists())

    def seed_private_trace(self, quotation):
        """The publication trace (S2) is supplier-only: stock at quote time, sources and acknowledgements."""
        DealQuotationAudit.objects.filter(quotation_id=quotation['id']).update(settings_snapshot={'note': 'SECRETO-AUDITORIA'})
        DealQuotationLineAudit.objects.filter(line__quotation_id=quotation['id']).update(explanation={'list_price': '777.77', 'rule': 'SECRETO-REGLA'})
        self.assertEqual(DealQuotationLineAudit.objects.filter(line__quotation_id=quotation['id']).count(), 2)

    def pricing_paths(self, account):
        """The supplier's private pricing endpoints, plus the same endpoints under the caller's own account."""
        private, own = f'/api/v1/accounts/{self.supplier_a.pk}', f'/api/v1/accounts/{account.pk}'
        return [f'{private}/pricing/settings/', f'{private}/price-lists/', f'{private}/prices/', f'{private}/prices/?status=below_floor',
                f'{private}/prices/{self.item_a.pk}/history/', f'{private}/prices/export/', f'{private}/pricing/history/',
                f'{own}/price-lists/', f'{own}/prices/', f'{own}/prices/{self.item_a.pk}/history/', f'{own}/prices/export/', f'{own}/pricing/history/',
                f'{private}/prices/import/template/', f'{private}/prices/import/jobs/{self.import_job}/', f'{private}/prices/import/jobs/{self.import_job}/errors/',
                f'{own}/prices/import/template/', f'{own}/prices/import/jobs/{self.import_job}/', f'{own}/prices/import/jobs/{self.import_job}/errors/',
                f'{private}/clients/', f'{private}/clients/?search=SECRETO', f'{private}/clients/{self.client_account.pk}/profile/', f'{private}/pricing-rules/',
                f'{private}/pricing-rules/?client={self.client_account.pk}', f'{own}/clients/', f'{own}/clients/{self.client_account.pk}/profile/',
                f'{own}/pricing-rules/', f'{own}/pricing-rules/?client={self.client_account.pk}']

    def calls(self, user, account, order, *, writes=()):
        self.client.force_authenticate(user)
        account_url, part = f'/api/v1/accounts/{account.pk}', self.part.pk
        responses = {path: self.client.get(path) for path in [
            '/api/v1/accounts/', '/api/v1/catalog/?include_facets=1', f'/api/v1/catalog/?search={self.part.sku}&include_facets=1',
            f'/api/v1/catalog/{part}/technical/', f'/api/v1/catalog/{part}/suppliers/', '/api/v1/wishlist/', '/api/v1/wishlist/state/',
            f'{account_url}/catalog/{part}/request-state/', f'{account_url}/sent-requests/', f'{account_url}/sent-requests/{order.pk}/',
            f'{account_url}/requests/', f'{account_url}/requests/{order.pk}/', self.messages(order, account), *self.pricing_paths(account), *self.trace_paths]}
        responses[f'PUT /api/v1/wishlist/{part}/'] = self.client.put(f'/api/v1/wishlist/{part}/')
        responses['POST messages'] = self.client.post(self.messages(order, account), {'message_id': str(uuid.uuid4()), 'body': 'HOLA'}, format='json')
        responses['POST analytics'] = self.client.post('/api/v1/analytics/events/', {
            'event_id': str(uuid.uuid4()), 'visit_id': str(uuid.uuid4()), 'account_id': str(account.pk), 'kind': 'basket_add', 'part_id': str(part),
            'supplier_item_id': str(self.item_a.pk), 'quantity': 1}, format='json')
        private = f'/api/v1/accounts/{self.supplier_a.pk}'
        responses['POST prices'] = self.client.post(f'{private}/prices/', {'changes': [{'list_id': str(PriceList.objects.get().pk), 'supplier_item_id': str(self.item_a.pk),
                                                                                         'unit_price': '1.00', 'expected_revision': 1}]}, format='json')
        responses['POST price-lists'] = self.client.post(f'{private}/price-lists/', {'code': 'ROBADA', 'name': 'ROBADA'}, format='json')
        responses['POST prices/import'] = self.client.post(f'{private}/prices/import/', {'file': self.price_file([['A-001', 'ROBADA', '', '1.00']])}, format='multipart')
        responses['POST prices/import (propia)'] = self.client.post(f'{account_url}/prices/import/', {'file': self.price_file([['A-001', '', '', '1.00']])}, format='multipart')
        responses['POST import commit'] = self.client.post(f'{private}/prices/import/jobs/{self.import_job}/', {'batch_index': 0}, format='json')
        client, item = str(self.client_account.pk), [{'supplier_item_id': str(self.item_a.pk), 'quantity': 2}]
        responses['POST profile'] = self.client.post(f'{private}/clients/{client}/profile/', {'operation_id': str(uuid.uuid4()), 'expected_version': 1,
                                                                                             'internal_notes': 'ROBADA'}, format='json')
        responses['POST pricing-rules'] = self.client.post(f'{private}/pricing-rules/', {'operation_id': str(uuid.uuid4()), 'name': 'ROBADA', 'scope': 'all',
                                                                                        'target': 'all', 'kind': 'discount', 'value': '90'}, format='json')
        responses['POST rule update'] = self.client.post(f'{private}/pricing-rules/{self.rule.pk}/', {'expected_version': 1, 'value': '1.00'}, format='json')
        responses['POST rule archive'] = self.client.post(f'{private}/pricing-rules/{self.rule.pk}/archive/', {'expected_version': 1}, format='json')
        responses['POST simulate'] = self.client.post(f'{private}/pricing/simulate/', {'client_id': client, 'items': item}, format='json')
        responses['POST simulate (propio)'] = self.client.post(f'{account_url}/pricing/simulate/', {'client_id': client, 'items': item}, format='json')
        for name, call in writes:
            responses[name] = call()
        return responses

    def oratek_calls(self):
        """Oratek superusers have no supplier membership: management and analytics never show supplier prices (decision D5)."""
        self.client.force_authenticate(self.root)
        return {path: self.client.get(path) for path in [
            '/api/v1/management/inventory/', '/api/v1/management/catalog/', '/api/v1/management/analytics/', f'/api/v1/management/accounts/{self.supplier_a.pk}/',
            f'/api/v1/management/inventory/{self.item_a.pk}/', *self.pricing_paths(self.supplier_a)]}

    def assert_no_markers(self, label, responses):
        for path, response in responses.items():
            self.assertLess(response.status_code, 500, f'{label} {path}')
            if response.get('Content-Type') == XLSX_TYPE:
                body = json.dumps([[cell.value for cell in row] for sheet in load_workbook(io.BytesIO(response.content)) for row in sheet.iter_rows()], default=str)
            else:
                body = response.content.decode()
            for marker in MARKERS:
                self.assertNotIn(marker, body, f'{label} {path} filtró {marker}')

    def test_private_pricing_markers_never_reach_clients_or_competitors(self):
        capture, root = LogCapture(), logging.getLogger()
        level = root.level
        root.addHandler(capture)
        root.setLevel(logging.INFO)
        try:
            self.canary()
        finally:
            root.removeHandler(capture)
            root.setLevel(level)
        for marker in MARKERS:
            self.assertNotIn(marker, '\n'.join(capture.lines), f'Los registros filtraron {marker}')

    def canary(self):
        order = self.order()
        self.seed_private_markers(order)
        # The supplier itself does see its private data (so the markers are real and reachable).
        self.client.force_authenticate(self.seller_a)
        own = {path: self.client.get(path).content.decode() for path in self.pricing_paths(self.supplier_a)[1:4]}
        self.assertTrue(all(marker in ''.join(own.values()) for marker in ['777.77', '666.66', 'SECRETO-LINEA', 'SECRETA_X']))
        own = ''.join(self.client.get(path).content.decode() for path in self.pricing_paths(self.supplier_a) if '/clients/' in path or '/pricing-rules/' in path)
        self.assertTrue(all(marker in own for marker in ['555.55', 'SECRETO-REGLA', 'SECRETO-PERFIL', 'SECRETO-ERP', 'SECRETO-NOTA']))
        self.review(order)
        quote = self.quote(order).data['quotation']
        self.seed_private_trace(quote)
        self.trace_paths = [f'/api/v1/accounts/{account.pk}/requests/{order.pk}/quotations/{quote["id"]}/trace/'
                            for account in (self.supplier_a, self.client_account, self.supplier_b)]
        responses = self.calls(self.buyer, self.client_account, order, writes=[
            ('POST request_adjustment', lambda: self.action(order, 'request_adjustment', quotation_id=quote['id'], reason='MENOS UNIDADES'))])
        self.assertIn('"unit_price":"13.00"', responses[f'/api/v1/accounts/{self.client_account.pk}/sent-requests/{order.pk}/'].content.decode())
        self.seed_private_draft(order)
        draft_path = f'/requests/{order.pk}/draft/'
        drafting = self.calls(self.buyer, self.client_account, order, writes=[
            ('GET borrador del proveedor', lambda: self.client.get(f'/api/v1/accounts/{self.supplier_a.pk}{draft_path}')),
            ('GET borrador como cliente', lambda: self.client.get(f'/api/v1/accounts/{self.client_account.pk}{draft_path}'))])
        self.assertEqual((drafting['GET borrador del proveedor'].status_code, drafting['GET borrador como cliente'].status_code), (404, 403))
        responses.update({f'{path} (con borrador)': response for path, response in drafting.items()})
        self.assert_no_markers('proveedor B con borrador', self.calls(self.seller_b, self.supplier_b, order, writes=[
            ('GET borrador ajeno', lambda: self.client.get(f'/api/v1/accounts/{self.supplier_a.pk}{draft_path}')),
            ('GET borrador desde su cuenta', lambda: self.client.get(f'/api/v1/accounts/{self.supplier_b.pk}{draft_path}'))]))
        self.client.force_authenticate(self.seller_a)
        self.assertEqual(self.action(order, 'return_quote', quotation_id=quote['id']).status_code, 200)
        self.assertFalse(DealQuotationDraft.objects.exists())
        self.client.force_authenticate(self.buyer)
        # A stock drop blocks the accept: the client learns nothing about quantities, prices or the trace.
        SupplierItem.objects.filter(pk=self.item_a.pk).update(reported_quantity=3)
        SupplierPricingSettings.objects.filter(supplier=self.supplier_a).update(accept_shortfall_policy='block')
        responses['POST accept bloqueado'] = blocked = self.action(order, 'accept', quotation_id=quote['id'])
        self.assertEqual((blocked.status_code, set(blocked.data)), (409, {'detail'}))
        self.assertNotRegex(blocked.content.decode(), r'\d')
        SupplierItem.objects.filter(pk=self.item_a.pk).update(reported_quantity=10)
        responses['POST accept'] = self.action(order, 'accept', quotation_id=quote['id'])
        responses['POST requests'] = self.submit(self.payload([(self.item_a, 1), (self.item_b, 1)]))
        self.assertEqual((responses['POST accept'].status_code, responses['POST requests'].status_code), (200, 200))
        self.assert_no_markers('cliente A', responses)
        buyer_b = User.objects.create_user('privacy-buyer-b', 'privacy-buyer-b@example.invalid', 'Request938!')
        client_b = self.account('CLIENTE B', 'client_walkin', buyer_b)
        self.assert_no_markers('cliente B', self.calls(buyer_b, client_b, order, writes=[
            ('POST accept', lambda: self.action(order, 'accept', account=client_b, quotation_id=quote['id'])),
            ('POST requests', lambda: self.submit(self.payload([(self.item_a, 2)]), account=client_b))]))
        self.assert_no_markers('proveedor B', self.calls(self.seller_b, self.supplier_b, order, writes=[
            ('POST quote', lambda: self.action(order, 'quote', account=self.supplier_b, lines=[]))]))
        self.assert_no_markers('Oratek', self.oratek_calls())
        self.assertEqual(DealCommand.objects.filter(account=self.client_account).count(), 2)
        for command in DealCommand.objects.all():
            for marker in MARKERS:
                self.assertNotIn(marker, json.dumps(command.result))
        self.assertEqual((PriceListEntry.objects.get(item=self.item_a).unit_price, PriceList.objects.count()), (Decimal('777.77'), 1))
        # Nobody else edits the pair profile or the supplier's rules.
        self.assertEqual(list(ClientPricingProfile.objects.filter(supplier=self.supplier_a).values_list('internal_notes', 'version')), [('SECRETO-PERFIL', 1)])
        self.assertEqual(sorted(PricingRule.objects.filter(supplier=self.supplier_a).values_list('name', 'revision', 'active')),
                         [('SECRETO-REGLA', 1, True), ('SECRETO-REGLA GENERAL', 1, True)])
        self.assertFalse(PricingRule.objects.filter(name='ROBADA').exists())
        # Nobody else uploads into or applies the supplier's import; the competitor's own uploads stay in its own account.
        self.assertEqual(list(PriceImportJob.objects.filter(supplier=self.supplier_a).values_list('owner', 'completed_batches')), [(self.seller_a.pk, 0)])
        self.assertEqual(set(PriceImportJob.objects.values_list('supplier', flat=True)), {self.supplier_a.pk, self.supplier_b.pk})

    def test_client_facing_payload_key_sets_are_pinned(self):
        receipt = self.submit(self.payload([(self.item_a, 3), (self.item_a2, 2)]))
        self.assertEqual(set(receipt.data), {'submission_id', 'created_at', 'requests'})
        self.assertEqual(set(receipt.data['requests'][0]), {'id', 'reference', 'supplier', 'line_count', 'unit_count', 'added_unit_count', 'appended'})
        order = SupplierRequest.objects.get(pk=receipt.data['requests'][0]['id'])
        self.review(order)
        quoted = self.quote(order)
        self.assertEqual(set(quoted.data), SUMMARY_KEYS | DEAL_KEYS | {'notes', 'lines'})
        self.assertEqual({key for line in quoted.data['lines'] for key in line}, SUPPLIER_LINE_KEYS)
        self.client.force_authenticate(self.seller_a)
        listed = self.client.get(self.url(self.supplier_a)).data['results'][0]
        self.assertEqual(set(listed), SUMMARY_KEYS | {'client', 'draft_state', 'availability_alert'})
        self.assertEqual((listed['draft_state'], listed['availability_alert']), (None, None))
        self.assertEqual(set(self.client.get(f'{self.url(self.supplier_a)}{order.pk}/').data), set(quoted.data))
        self.client.force_authenticate(self.buyer)
        base = f'/api/v1/accounts/{self.client_account.pk}'
        sent = self.client.get(f'{base}/sent-requests/').data['results'][0]
        self.assertEqual(set(sent), SENT_SUMMARY_KEYS)
        detail = self.client.get(f'{base}/sent-requests/{order.pk}/').data
        accepted = self.action(order, 'accept', quotation_id=detail['quotation']['id']).data
        for payload in [detail, accepted, DealCommand.objects.get(account=self.client_account).result]:
            self.assertEqual(set(payload), SENT_SUMMARY_KEYS | DEAL_KEYS | {'notes', 'lines'})
            self.assertEqual({key for line in payload['lines'] for key in line}, CLIENT_LINE_KEYS)
            self.assertEqual(set(payload['quotation']), QUOTATION_KEYS)
            self.assertEqual({key for quotation in payload['quotations'] for key in quotation}, QUOTATION_KEYS)
            self.assertEqual({key for line in payload['quotation']['lines'] for key in line}, QUOTATION_LINE_KEYS)
            self.assertEqual({key for event in payload['events'] for key in event}, EVENT_KEYS)
            self.assertEqual((set(payload['client']), set(payload['supplier'])), ({'id', 'name'}, {'id', 'name'}))
        self.assertEqual(set(DealCommand.objects.get(account=self.supplier_a).result), SUMMARY_KEYS | DEAL_KEYS | {'notes', 'lines'})
        state = self.client.get(f'{base}/catalog/{self.part.pk}/request-state/').data
        self.assertEqual((set(state), set(state['items'][0])), ({'part_id', 'totals', 'items'}, {
            'supplier_item_id', 'supplier_id', 'codigo', 'brand', 'sent_quantity', 'pending_quantity', 'reviewed_quantity',
            'quoted_quantity', 'adjustment_quantity', 'handshaked_quantity'}))
        message = self.client.post(self.messages(order), {'message_id': str(uuid.uuid4()), 'body': 'GRACIAS'}, format='json').data
        self.assertEqual(set(message), {'id', 'message_id', 'body', 'account_id', 'account_name', 'actor_name', 'created_at'})
        self.assertEqual(set(self.client.get('/api/v1/accounts/').data['results'][0]), {'id', 'name', 'roles', 'capabilities', 'permission'})
        part = self.client.get('/api/v1/catalog/').data['results'][0]
        self.assertEqual(set(part), {'id', 'sku', 'is_OEM', 'name', 'description', 'category', 'subcategory', 'part_type', 'codes', 'availability', 'images'})
        self.assertEqual(set(part['availability']), {'status', 'supplier_count', 'updated_at'})
        offers = self.client.get(f'/api/v1/catalog/{self.part.pk}/suppliers/').data
        self.assertEqual({key for offer in offers for key in offer}, {'supplier_id', 'supplier_name', 'in_stock', 'items'})
        self.assertEqual({key for offer in offers for item in offer['items'] for key in item}, {'id', 'codigo', 'brand', 'description'})

    def test_pricing_models_stay_out_of_admin_and_management_serializers(self):
        for model in [SupplierPricingSettings, PricingAuditEvent, DealQuotationDraft, DealQuotationDraftLine, DealQuotationAudit, DealQuotationLineAudit,
                      PriceList, PriceListEntry, SupplierItemPricing, PriceChange, PriceImportJob, PriceImportBatch, ClientPricingProfile, PricingRule]:
            self.assertNotIn(model, admin.site._registry)
        self.assertEqual(SupplierItemSerializer.Meta.fields, ['id', 'supplier_invent_id', 'part', 'codigo', 'brand', 'description', 'references',
                                                              'matching_status', 'source', 'reported_quantity', 'reserved_quantity',
                                                              'available_quantity', 'updated_at'])
        self.assertEqual(ManagedInventorySerializer.Meta.fields, SupplierItemSerializer.Meta.fields + ['supplier', 'supplier_name', 'part_name', 'part_sku'])
