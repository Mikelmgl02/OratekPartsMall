"""Supplier-only pricing endpoints. Nothing here feeds catalog, offers, request or deal payloads, and nothing is reachable by
Oratek staff (no admin, no management endpoint): only members of the supplier account read or write these rows."""
import io
import re
import uuid
from decimal import Decimal

from django.db import IntegrityError, transaction
from django.db.models import Count, Exists, F, OuterRef, Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from drf_spectacular.types import OpenApiTypes
from drf_spectacular.utils import OpenApiParameter, extend_schema
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter
from rest_framework import serializers
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.pagination import PageNumberPagination
from rest_framework.response import Response
from rest_framework.views import APIView

from .models import Account, SupplierItem
from .pricing_models import (CURRENCY_CHOICES, OWNER_ONLY_SETTINGS, PERMISSION_CHOICES, POLICY_CHOICES, PREFILL_QUANTITY_CHOICES, PRICE_CHANGE_KINDS,
                             PRICE_LIST_CODE, PRICE_WRITE_SOURCES, SHORTFALL_POLICY_CHOICES, PriceChange, PriceList, PriceListEntry, PricingAuditEvent,
                             SupplierItemPricing, choice_values, pricing_settings, record_pricing_event)
from .pricing_services import PriceConflict, set_prices
from .views import PERMISSION_RANK, membership_for

SETTINGS_FIELDS = ('default_currency', 'usd_pab_parity', 'config_min_permission', 'publish_min_permission', 'over_request_policy',
                   'over_stock_policy', 'accept_shortfall_policy', 'prefill_quantity', 'assistant_enabled')


class PricingSettingsFields(serializers.Serializer):
    default_currency = serializers.ChoiceField(choices=choice_values(CURRENCY_CHOICES))
    usd_pab_parity = serializers.BooleanField()
    config_min_permission = serializers.ChoiceField(choices=PERMISSION_CHOICES)
    publish_min_permission = serializers.ChoiceField(choices=PERMISSION_CHOICES)
    over_request_policy = serializers.ChoiceField(choices=choice_values(POLICY_CHOICES))
    over_stock_policy = serializers.ChoiceField(choices=choice_values(POLICY_CHOICES))
    accept_shortfall_policy = serializers.ChoiceField(choices=choice_values(SHORTFALL_POLICY_CHOICES))
    prefill_quantity = serializers.ChoiceField(choices=choice_values(PREFILL_QUANTITY_CHOICES))
    assistant_enabled = serializers.BooleanField()


class MemberName(serializers.Serializer):
    name = serializers.CharField()


class PricingSettingsSerializer(PricingSettingsFields):
    version = serializers.IntegerField()
    updated_at = serializers.DateTimeField(allow_null=True)
    updated_by = MemberName(allow_null=True)
    permission = serializers.ChoiceField(choices=PERMISSION_CHOICES, help_text='Tu permiso en la cuenta.')
    can_configure = serializers.BooleanField()
    can_manage_permissions = serializers.BooleanField()


class PricingSettingsUpdate(PricingSettingsFields):
    expected_version = serializers.IntegerField(min_value=1)

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        for field in SETTINGS_FIELDS:
            self.fields[field].required = False


class PricingSettingsConflict(serializers.Serializer):
    detail = serializers.CharField()
    settings = PricingSettingsSerializer()


def settings_data(row, membership):
    rank = PERMISSION_RANK.get(membership.permission, -1)
    user = row.updated_by
    return PricingSettingsSerializer({**{field: getattr(row, field) for field in SETTINGS_FIELDS}, 'version': row.version,
        'updated_at': row.updated_at, 'updated_by': {'name': user.get_full_name() or user.username} if user else None,
        'permission': membership.permission, 'can_configure': rank >= PERMISSION_RANK[row.config_min_permission],
        'can_manage_permissions': membership.permission == 'owner'}).data


class PricingSettingsView(APIView):
    @extend_schema(operation_id='v1_accounts_pricing_settings_retrieve', responses=PricingSettingsSerializer,
                   description='Configuración privada de precios del proveedor. Leerla no crea registros.')
    def get(self, request, account_id):
        account, membership = membership_for(request.user, account_id, 'supplier')
        return Response(settings_data(pricing_settings(account), membership))

    @extend_schema(operation_id='v1_accounts_pricing_settings_update', request=PricingSettingsUpdate,
                   responses={200: PricingSettingsSerializer, 409: PricingSettingsConflict},
                   description='Actualización parcial con expected_version. Los permisos y el asistente de IA solo los cambia el propietario.')
    @transaction.atomic
    def post(self, request, account_id):
        account, membership = membership_for(request.user, account_id, 'supplier')
        row = pricing_settings(account, lock=True)
        if PERMISSION_RANK.get(membership.permission, -1) < PERMISSION_RANK[row.config_min_permission]:
            raise PermissionDenied('Tu permiso en la cuenta no permite cambiar esta configuración. Pide a un administrador de tu cuenta que la cambie.')
        serializer = PricingSettingsUpdate(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        changes = {field: data[field] for field in SETTINGS_FIELDS if field in data and getattr(row, field) != data[field]}
        if membership.permission != 'owner' and set(changes) & set(OWNER_ONLY_SETTINGS):
            raise PermissionDenied('Solo el propietario de la cuenta puede cambiar quién configura precios, quién envía cotizaciones y el asistente de IA.')
        # Values already in place are a no-op, so a retry after a lost response never conflicts.
        if changes and row.version != data['expected_version']:
            return Response({'detail': 'Otro miembro de tu equipo cambió esta configuración. Revisa los valores actuales.',
                             'settings': settings_data(row, membership)}, status=409)
        if changes:
            old = {field: getattr(row, field) for field in changes}
            for field, value in changes.items():
                setattr(row, field, value)
            row.version += 1
            row.updated_by = request.user
            row.save()
            record_pricing_event(account, request.user, 'settings_changed', payload={'version': row.version, 'old': old, 'new': changes})
        return Response(settings_data(row, membership))


# ---------------------------------------------------------------------------------------------------------------------------
# Price lists, manual price editing, history and export (S3).

CONFIG_DENIED = 'Tu permiso en la cuenta no permite cambiar precios. Pide a un administrador de tu cuenta que lo haga.'
XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
# PRECIO_MINIMO and PRECIO_PISO are the floor-price columns of the import and export files, so no list may own those headers.
RESERVED_LIST_CODES = ('MINIMO', 'PISO')


def can_configure(account, membership, settings_row=None):
    settings_row = settings_row or pricing_settings(account)
    return PERMISSION_RANK.get(membership.permission, -1) >= PERMISSION_RANK[settings_row.config_min_permission]


def configurer(user, account_id):
    """Members at or above config_min_permission (staff by default) may write prices and lists."""
    account, membership = membership_for(user, account_id, 'supplier')
    if not can_configure(account, membership):
        raise PermissionDenied(CONFIG_DENIED)
    return account, membership


def user_name(user):
    return {'name': user.get_full_name() or user.username} if user else None


class PricePagination(PageNumberPagination):
    page_size = 100


def page_serializer(name, item):
    return type(name, (serializers.Serializer,), {'count': serializers.IntegerField(), 'next': serializers.CharField(allow_null=True),
                                                  'previous': serializers.CharField(allow_null=True), 'results': item(many=True)})


class PriceListSerializer(serializers.Serializer):
    id = serializers.UUIDField()
    code = serializers.CharField()
    name = serializers.CharField()
    currency = serializers.ChoiceField(choices=choice_values(CURRENCY_CHOICES))
    is_default = serializers.BooleanField()
    active = serializers.BooleanField()
    version = serializers.IntegerField()
    priced_count = serializers.IntegerField(help_text='Artículos con precio en esta lista.')
    missing_count = serializers.IntegerField(help_text='Artículos de tu inventario sin precio en esta lista.')
    created_at = serializers.DateTimeField()
    updated_at = serializers.DateTimeField()


class PriceListsEnvelope(serializers.Serializer):
    can_configure = serializers.BooleanField()
    item_count = serializers.IntegerField()
    results = PriceListSerializer(many=True)


class PriceListFields(serializers.Serializer):
    name = serializers.CharField(max_length=120)
    currency = serializers.ChoiceField(choices=choice_values(CURRENCY_CHOICES))
    is_default = serializers.BooleanField()

    def validate_name(self, value):
        value = value.strip().upper()
        if not value:
            raise serializers.ValidationError('Escribe el nombre de la lista.')
        return value


class PriceListCreate(PriceListFields):
    code = serializers.CharField(max_length=30)

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields['currency'].required = self.fields['is_default'].required = False

    def validate_code(self, value):
        value = value.strip().upper()
        if not re.fullmatch(PRICE_LIST_CODE, value):
            raise serializers.ValidationError('Usa de 1 a 30 letras sin tilde, números o guiones bajos (A-Z, 0-9, _).')
        if value in RESERVED_LIST_CODES:
            raise serializers.ValidationError(f'El código {value} está reservado para el precio mínimo (PRECIO_{value}); elige otro.')
        return value


class PriceListUpdate(PriceListFields):
    expected_version = serializers.IntegerField(min_value=1)
    active = serializers.BooleanField()

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        for field in ('name', 'currency', 'is_default', 'active'):
            self.fields[field].required = False


class PriceListConflict(serializers.Serializer):
    detail = serializers.CharField()
    price_list = PriceListSerializer()


def price_list_rows(account, lists=None):
    rows = lists if lists is not None else PriceList.objects.filter(supplier=account)
    items = SupplierItem.objects.filter(supplier=account).count()
    return [{**{field: getattr(row, field) for field in ('id', 'code', 'name', 'currency', 'is_default', 'active', 'version', 'created_at', 'updated_at')},
             'priced_count': row.priced, 'missing_count': max(0, items - row.priced)}
            for row in rows.annotate(priced=Count('entries')).order_by('-is_default', 'code')], items


def demote(account, keep=None):
    # The previous default changes too: its version moves, so an edit based on what it was before is refused (409), not applied.
    PriceList.objects.filter(supplier=account, is_default=True).exclude(pk=getattr(keep, 'pk', None)).update(
        is_default=False, version=F('version') + 1, updated_at=timezone.now())


def price_list_data(row):
    return PriceListSerializer(price_list_rows(row.supplier, PriceList.objects.filter(pk=row.pk))[0][0]).data


class PriceListsView(APIView):
    @extend_schema(operation_id='v1_accounts_price_lists_list', responses=PriceListsEnvelope,
                   description='Listas de precios privadas del proveedor, con artículos con y sin precio.')
    def get(self, request, account_id):
        account, membership = membership_for(request.user, account_id, 'supplier')
        rows, items = price_list_rows(account)
        return Response(PriceListsEnvelope({'can_configure': can_configure(account, membership), 'item_count': items, 'results': rows}).data)

    @extend_schema(operation_id='v1_accounts_price_lists_create', request=PriceListCreate, responses={200: PriceListSerializer, 201: PriceListSerializer},
                   description='Crea una lista. La primera lista del proveedor queda como predeterminada. Repetir la misma creación devuelve la lista.')
    @transaction.atomic
    def post(self, request, account_id):
        account, _ = configurer(request.user, account_id)
        serializer = PriceListCreate(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        Account.objects.select_for_update().get(pk=account.pk)
        currency = data.get('currency') or pricing_settings(account).default_currency
        existing = PriceList.objects.filter(supplier=account, code=data['code']).first()
        if existing:
            # A retry after a lost response finds the list it created; any other reuse of the code is refused.
            if existing.name == data['name'] and existing.currency == currency:
                return Response(price_list_data(existing))
            raise ValidationError({'code': [f'Ya tienes una lista con el código {data["code"]}.']})
        first = not PriceList.objects.filter(supplier=account).exists()
        default = first or data.get('is_default', False)
        if default:
            demote(account)
        try:
            with transaction.atomic():
                row = PriceList.objects.create(supplier=account, code=data['code'], name=data['name'], currency=currency, is_default=default, created_by=request.user)
        except IntegrityError:
            raise ValidationError('La lista no cumple las reglas de listas de precios. Actualiza la página e inténtalo de nuevo.')
        record_pricing_event(account, request.user, 'price_list_changed', object_id=str(row.pk), payload={
            'action': 'created', 'code': row.code, 'new': {'name': row.name, 'currency': row.currency, 'is_default': row.is_default}})
        return Response(price_list_data(row), status=201)


class PriceListDetail(APIView):
    @extend_schema(operation_id='v1_accounts_price_lists_update', request=PriceListUpdate, responses={200: PriceListSerializer, 409: PriceListConflict},
                   description='Cambia nombre, moneda (solo sin precios), predeterminada o activa, con expected_version. Valores iguales no cambian nada.')
    @transaction.atomic
    def post(self, request, account_id, pk):
        account, _ = configurer(request.user, account_id)
        serializer = PriceListUpdate(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        Account.objects.select_for_update().get(pk=account.pk)
        row = get_object_or_404(PriceList.objects.select_for_update().filter(supplier=account), pk=pk)
        changes = {field: data[field] for field in ('name', 'currency', 'is_default', 'active') if field in data and getattr(row, field) != data[field]}
        if not changes:
            return Response(price_list_data(row))
        if row.version != data['expected_version']:
            return Response({'detail': 'Otro usuario actualizó esta lista. Revisa sus valores actuales.', 'price_list': price_list_data(row)}, status=409)
        active, default = changes.get('active', row.active), changes.get('is_default', row.is_default)
        if 'currency' in changes and row.entries.exists():
            raise ValidationError({'currency': ['No puedes cambiar la moneda de una lista que ya tiene precios.']})
        if row.is_default and changes.get('is_default') is False:
            raise ValidationError({'is_default': ['Marca otra lista como predeterminada; siempre debe haber una.']})
        if default and not active:
            raise ValidationError({'active': ['Elige otra lista predeterminada antes de archivar esta.' if row.is_default
                                              else 'Activa la lista antes de marcarla como predeterminada.']})
        if changes.get('is_default'):
            demote(account, row)
        old = {field: getattr(row, field) for field in changes}
        for field, value in changes.items():
            setattr(row, field, value)
        row.version += 1
        try:
            with transaction.atomic():
                row.save()
        except IntegrityError:
            raise ValidationError('La lista no cumple las reglas de listas de precios. Actualiza la página e inténtalo de nuevo.')
        record_pricing_event(account, request.user, 'price_list_changed', object_id=str(row.pk), payload={'action': 'updated', 'code': row.code,
                                                                                                           'old': old, 'new': changes})
        return Response(price_list_data(row))


class PriceCell(serializers.Serializer):
    unit_price = serializers.DecimalField(max_digits=12, decimal_places=2)
    revision = serializers.IntegerField()
    updated_at = serializers.DateTimeField()


class PriceRow(serializers.Serializer):
    supplier_item_id = serializers.UUIDField()
    supplier_invent_id = serializers.CharField()
    codigo = serializers.CharField()
    brand = serializers.CharField(allow_blank=True)
    description = serializers.CharField(allow_blank=True)
    matching_status = serializers.ChoiceField(choices=SupplierItem._meta.get_field('matching_status').choices)
    discount_group = serializers.CharField(allow_blank=True, help_text='LINEA propia del proveedor.')
    floor_price = serializers.DecimalField(max_digits=12, decimal_places=2, allow_null=True)
    item_pricing_revision = serializers.IntegerField(allow_null=True)
    prices = serializers.DictField(child=PriceCell(), help_text='Precio por código de lista; una lista sin precio no aparece.')
    updated_at = serializers.DateTimeField(allow_null=True, help_text='Último cambio de precio, mínimo o línea.')


PricePage = page_serializer('PricePage', PriceRow)


class PriceFilters(serializers.Serializer):
    list = serializers.UUIDField(required=False)
    search = serializers.CharField(required=False, allow_blank=True, max_length=200)
    status = serializers.ChoiceField(required=False, choices=['priced', 'missing', 'below_floor'])


MONEY_ERRORS = {'invalid': 'Escribe un importe válido, por ejemplo 12.50.', 'max_decimal_places': 'Usa un máximo de dos decimales.',
                'max_digits': 'Usa como máximo 10 dígitos enteros.', 'max_whole_digits': 'Usa como máximo 10 dígitos enteros.'}


class PriceChangeInput(serializers.Serializer):
    list_id = serializers.UUIDField()
    supplier_item_id = serializers.UUIDField()
    unit_price = serializers.DecimalField(max_digits=12, decimal_places=2, min_value=Decimal('0.01'), allow_null=True, help_text='null elimina el precio.',
                                          error_messages={**MONEY_ERRORS, 'min_value': 'El precio de lista debe ser mayor que 0.'})
    expected_revision = serializers.IntegerField(min_value=1, allow_null=True, help_text='La revisión que viste; null si no tenía precio.')


class ItemPricingInput(serializers.Serializer):
    supplier_item_id = serializers.UUIDField()
    discount_group = serializers.CharField(max_length=60, allow_blank=True, required=False)
    floor_price = serializers.DecimalField(max_digits=12, decimal_places=2, min_value=Decimal('0'), allow_null=True, required=False,
                                           error_messages={**MONEY_ERRORS, 'min_value': 'El precio mínimo no puede ser negativo.'})
    expected_revision = serializers.IntegerField(min_value=1, allow_null=True)

    def validate_discount_group(self, value):
        return value.strip().upper()

    def validate(self, data):
        if 'discount_group' not in data and 'floor_price' not in data:
            raise serializers.ValidationError('Indica la línea o el precio mínimo.')
        return data


class PriceWrite(serializers.Serializer):
    changes = PriceChangeInput(many=True, required=False, max_length=500)
    items = ItemPricingInput(many=True, required=False, max_length=500)

    def validate(self, data):
        changes, items = data.get('changes', []), data.get('items', [])
        if not changes and not items:
            raise serializers.ValidationError('No hay cambios de precio para guardar.')
        if len({(row['list_id'], row['supplier_item_id']) for row in changes}) != len(changes) or len({row['supplier_item_id'] for row in items}) != len(items):
            raise serializers.ValidationError('Incluye cada precio una sola vez en cada guardado.')
        return data


class PriceWriteSummary(serializers.Serializer):
    created = serializers.IntegerField()
    updated = serializers.IntegerField()
    removed = serializers.IntegerField()
    unchanged = serializers.IntegerField()
    floor_updates = serializers.IntegerField()
    line_updates = serializers.IntegerField()


class PriceWriteResult(serializers.Serializer):
    summary = PriceWriteSummary()
    rows = PriceRow(many=True, help_text='Valores actuales de los artículos incluidos.')


class PriceConflictItem(serializers.Serializer):
    supplier_item_id = serializers.UUIDField()
    list_id = serializers.UUIDField(allow_null=True, help_text='null para línea o precio mínimo.')
    current = serializers.DictField()


class PriceWriteConflict(serializers.Serializer):
    detail = serializers.CharField()
    conflicts = PriceConflictItem(many=True)
    rows = PriceRow(many=True)


def item_pricing(item):
    try:
        return item.pricing
    except SupplierItem.pricing.RelatedObjectDoesNotExist:
        return None


def price_rows(account, items):
    """Grid rows for these items: every list's price keyed by code, plus LINEA, floor and their revisions."""
    items = list(items)
    codes = dict(PriceList.objects.filter(supplier=account).values_list('pk', 'code'))
    prices = {}
    for entry in PriceListEntry.objects.filter(price_list_id__in=codes, item__in=items):
        prices.setdefault(entry.item_id, {})[codes[entry.price_list_id]] = entry
    rows = []
    for item in items:
        pricing, cells = item_pricing(item), prices.get(item.pk, {})
        stamps = [entry.updated_at for entry in cells.values()] + ([pricing.updated_at] if pricing else [])
        rows.append({'supplier_item_id': item.pk, 'supplier_invent_id': item.supplier_invent_id, 'codigo': item.codigo, 'brand': item.brand,
                     'description': item.description, 'matching_status': item.matching_status, 'discount_group': pricing.discount_group if pricing else '',
                     'floor_price': pricing.floor_price if pricing else None, 'item_pricing_revision': pricing.revision if pricing else None,
                     'prices': {code: {'unit_price': entry.unit_price, 'revision': entry.revision, 'updated_at': entry.updated_at} for code, entry in cells.items()},
                     'updated_at': max(stamps) if stamps else None})
    return PriceRow(rows, many=True).data


def supplier_items(account):
    return SupplierItem.objects.filter(supplier=account).select_related('pricing').order_by('supplier_invent_id', 'id')


class PricesView(APIView):
    @extend_schema(operation_id='v1_accounts_prices_list', responses=PricePage, parameters=[
        OpenApiParameter('list', OpenApiTypes.UUID, description='Lista para filtrar por estado (predeterminada si se omite).'),
        OpenApiParameter('search', OpenApiTypes.STR), OpenApiParameter('status', OpenApiTypes.STR, enum=['priced', 'missing', 'below_floor']),
        OpenApiParameter('page', OpenApiTypes.INT)], description='Una fila por artículo de tu inventario, vinculado o no, con sus precios privados.')
    def get(self, request, account_id):
        account, _ = membership_for(request.user, account_id, 'supplier')
        filters = PriceFilters(data=request.query_params)
        filters.is_valid(raise_exception=True)
        data = filters.validated_data
        rows = supplier_items(account)
        if search := data.get('search', '').strip():
            rows = rows.filter(Q(supplier_invent_id__icontains=search) | Q(codigo__icontains=search) | Q(brand__icontains=search)
                               | Q(description__icontains=search) | Q(pricing__discount_group__icontains=search))
        target = get_object_or_404(PriceList.objects.filter(supplier=account), pk=data['list']) if 'list' in data else \
            PriceList.objects.filter(supplier=account, is_default=True).first()
        if data.get('status'):
            entries = PriceListEntry.objects.filter(price_list=target, item=OuterRef('pk'))
            rows = rows.filter(Exists(entries)) if data['status'] == 'priced' else rows.exclude(Exists(entries)) if data['status'] == 'missing' \
                else rows.filter(Exists(entries.filter(unit_price__lt=OuterRef('pricing__floor_price'))))
        paginator = PricePagination()
        page = paginator.paginate_queryset(rows, request, view=self)
        return paginator.get_paginated_response(price_rows(account, page))

    @extend_schema(operation_id='v1_accounts_prices_update', request=PriceWrite, responses={200: PriceWriteResult, 409: PriceWriteConflict},
                   description='Edición masiva con revisión esperada, todo o nada. Una fila que ya tiene el valor pedido no cambia (reintento seguro).')
    def post(self, request, account_id):
        account, _ = configurer(request.user, account_id)
        serializer = PriceWrite(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        try:
            result = set_prices(account, request.user, changes=data.get('changes', []), items=data.get('items', []), source='manual',
                                reference=f'op:{uuid.uuid4()}')
        except PriceConflict as conflict:
            touched = {row['supplier_item_id'] for row in conflict.conflicts}
            return Response({'detail': 'Otro usuario cambió algunos precios.', 'conflicts': conflict.conflicts,
                             'rows': price_rows(account, supplier_items(account).filter(pk__in=touched))}, status=409)
        return Response({'summary': result['summary'], 'rows': price_rows(account, supplier_items(account).filter(pk__in=result['item_ids']))})


class PriceHistoryEntry(serializers.Serializer):
    id = serializers.IntegerField()
    kind = serializers.ChoiceField(choices=choice_values(PRICE_CHANGE_KINDS))
    price_list = serializers.DictField(allow_null=True, help_text='{id, code} de la lista; null para línea o precio mínimo.')
    old_value = serializers.DecimalField(max_digits=12, decimal_places=2, allow_null=True)
    new_value = serializers.DecimalField(max_digits=12, decimal_places=2, allow_null=True)
    old_text = serializers.CharField(allow_blank=True)
    new_text = serializers.CharField(allow_blank=True)
    source = serializers.ChoiceField(choices=choice_values(PRICE_WRITE_SOURCES))
    reference = serializers.CharField(allow_blank=True)
    actor = MemberName()
    created_at = serializers.DateTimeField()


PriceHistoryPage = page_serializer('PriceHistoryPage', PriceHistoryEntry)


class PriceHistory(APIView):
    @extend_schema(operation_id='v1_accounts_prices_history_list', responses=PriceHistoryPage, parameters=[OpenApiParameter('page', OpenApiTypes.INT)],
                   description='Historial de precios, mínimo y línea de un artículo, del más reciente al más antiguo.')
    def get(self, request, account_id, item_id):
        account, _ = membership_for(request.user, account_id, 'supplier')
        item = get_object_or_404(SupplierItem.objects.filter(supplier=account), pk=item_id)
        rows = PriceChange.objects.filter(supplier=account, item=item).select_related('price_list', 'actor')
        paginator = PageNumberPagination()
        page = paginator.paginate_queryset(rows, request, view=self)
        return paginator.get_paginated_response(PriceHistoryEntry([{
            **{field: getattr(row, field) for field in ('id', 'kind', 'old_value', 'new_value', 'old_text', 'new_text', 'source', 'reference', 'created_at')},
            'price_list': {'id': str(row.price_list_id), 'code': row.price_list.code} if row.price_list_id else None,
            'actor': user_name(row.actor)} for row in page], many=True).data)


class PricingAuditEntry(serializers.Serializer):
    id = serializers.IntegerField()
    kind = serializers.CharField()
    actor = MemberName()
    client = serializers.DictField(allow_null=True, help_text='{id, name} del cliente, si aplica.')
    order = serializers.DictField(allow_null=True, help_text='{id, reference} de la orden, si aplica.')
    object_id = serializers.CharField(allow_blank=True)
    payload = serializers.DictField()
    created_at = serializers.DateTimeField()


PricingAuditPage = page_serializer('PricingAuditPage', PricingAuditEntry)


class PricingHistory(APIView):
    @extend_schema(operation_id='v1_accounts_pricing_history_list', responses=PricingAuditPage,
                   parameters=[OpenApiParameter('search', OpenApiTypes.STR), OpenApiParameter('page', OpenApiTypes.INT)],
                   description='Registro privado de cambios de precios, listas, borradores y cotizaciones de tu cuenta.')
    def get(self, request, account_id):
        account, _ = membership_for(request.user, account_id, 'supplier')
        search = str(request.query_params.get('search', '')).strip()[:200]
        rows = PricingAuditEvent.objects.filter(supplier=account).select_related('actor', 'client', 'order')
        if search:
            rows = rows.filter(Q(kind__icontains=search) | Q(object_id__icontains=search) | Q(order__reference__icontains=search)
                               | Q(client__name__icontains=search) | Q(actor__username__icontains=search))
        paginator = PageNumberPagination()
        page = paginator.paginate_queryset(rows, request, view=self)
        return paginator.get_paginated_response(PricingAuditEntry([{
            'id': row.pk, 'kind': row.kind, 'actor': user_name(row.actor), 'object_id': row.object_id, 'payload': row.payload, 'created_at': row.created_at,
            'client': {'id': str(row.client_id), 'name': row.client.name} if row.client_id else None,
            'order': {'id': str(row.order_id), 'reference': row.order.reference} if row.order_id else None} for row in page], many=True).data)


class PricesExport(APIView):
    @extend_schema(operation_id='v1_accounts_prices_export', responses={(200, XLSX_TYPE): bytes},
                   parameters=[OpenApiParameter('list', OpenApiTypes.UUID, description='Solo esta lista; sin ella, todas las listas activas.')],
                   description='Tus precios en el formato de importación: PRECIO es la lista predeterminada y PRECIO_<CÓDIGO> las demás.')
    def get(self, request, account_id):
        account, _ = membership_for(request.user, account_id, 'supplier')
        filters = PriceFilters(data=request.query_params)
        filters.is_valid(raise_exception=True)
        lists = PriceList.objects.filter(supplier=account)
        lists = list(lists.filter(pk=filters.validated_data['list']) if 'list' in filters.validated_data else lists.filter(active=True))
        if 'list' in filters.validated_data and not lists:
            raise ValidationError({'list': ['La lista de precios no existe.']})
        book = Workbook(write_only=True)
        sheet = book.create_sheet('PRECIOS')
        sheet.freeze_panes = 'A2'
        for index, width in enumerate([30, 26, 22, 60, 20, 16] + [16] * len(lists)):
            sheet.column_dimensions[get_column_letter(index + 1)].width = width
        sheet.append(['ID_INVENTARIO_PROVEEDOR', 'CODIGO', 'MARCA', 'DESCRIPCION', 'LINEA', 'PRECIO_MINIMO']
                     + ['PRECIO' if row.is_default else f'PRECIO_{row.code}' for row in lists])
        def cell(value, money=False):
            # Identifiers stay text (initial zeros, no formula); prices stay numbers with two decimals.
            result = WriteOnlyCell(sheet, value=value)
            result.number_format = '0.00' if money else '@'
            if not money and value is not None:
                result.data_type = 's'
            return result
        items = SupplierItem.objects.filter(supplier=account).select_related('pricing').order_by('supplier_invent_id', 'id')
        for offset in range(0, items.count(), 2000):
            chunk = list(items[offset:offset + 2000])
            prices = {}
            for entry in PriceListEntry.objects.filter(price_list__in=lists, item__in=chunk).values_list('item_id', 'price_list_id', 'unit_price'):
                prices[entry[:2]] = entry[2]
            for item in chunk:
                pricing = item_pricing(item)
                sheet.append([cell(item.supplier_invent_id), cell(item.codigo), cell(item.brand), cell(item.description),
                              cell(pricing.discount_group if pricing and pricing.discount_group else None),
                              cell(pricing.floor_price if pricing else None, money=True)]
                             + [cell(prices.get((item.pk, row.pk)), money=True) for row in lists])
        output = io.BytesIO()
        book.save(output)
        response = HttpResponse(output.getvalue(), content_type=XLSX_TYPE)
        response['Content-Disposition'] = 'attachment; filename="precios-proveedor.xlsx"'
        response['Cache-Control'] = 'private, no-store'
        return response
