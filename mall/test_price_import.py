import io
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta
from decimal import Decimal
from unittest import skipUnless
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.exceptions import ValidationError
from rest_framework.test import APITestCase, APITransactionTestCase

from . import test_deals as deal_tests
from . import test_requests as request_tests
from .matching_engine import MatchIndex, apply_item, part_row, reconcile_items
from .matching_worker import process_queue
from .matching_models import MatchingCase, MatchingQueue
from .models import InventoryUpdate, Membership, Part, StockEntry, SupplierItem
from .price_import import DELETE, PriceImportResult, parse_amount
from .price_import_models import PriceImportBatch, PriceImportJob
from .pricing_models import PriceChange, PriceList, PriceListEntry, PricingAuditEvent, SupplierItemPricing, SupplierPricingSettings
from .pricing_services import set_prices
from .services import ingest_inventory
from .supplier_import import stock_plan
from .supplier_import_models import SupplierInventoryImportJob

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


class Cell:
    """The attributes the cell readers use, for exact unit cases without building a workbook."""
    def __init__(self, value, data_type='n', is_date=False, number_format='General'):
        self.value, self.data_type, self.is_date, self.number_format = value, data_type, is_date, number_format


class PriceImportTests(APITestCase):
    setUp_workflow = request_tests.SupplierRequestWorkflowTests.setUp
    account = request_tests.SupplierRequestWorkflowTests.account
    item = request_tests.SupplierRequestWorkflowTests.item
    headers = ['ID_INVENTARIO_PROVEEDOR', 'CODIGO', 'PRECIO']

    def setUp(self):
        self.setUp_workflow()
        self.client.force_authenticate(self.seller_a)
        self.general = PriceList.objects.create(supplier=self.supplier_a, code='GENERAL', name='GENERAL', is_default=True, created_by=self.seller_a)

    def base(self, account=None):
        return f'/api/v1/accounts/{(account or self.supplier_a).pk}/prices/import/'

    def job_url(self, job, account=None):
        return f'{self.base(account)}jobs/{job}/'

    def file(self, rows, headers=None, title='PRECIOS', customize=None, name='PRECIOS-ERP.xlsx'):
        book = Workbook()
        sheet = book.active
        sheet.title = title
        sheet.append(headers or self.headers)
        for row in rows:
            sheet.append(row)
        if customize:
            customize(sheet)
        output = io.BytesIO()
        book.save(output)
        return SimpleUploadedFile(name, output.getvalue(), content_type=XLSX)

    def upload(self, rows, headers=None, user=None, account=None, **kwargs):
        if user:
            self.client.force_authenticate(user)
        return self.client.post(self.base(account), {'file': self.file(rows, headers, **kwargs)}, format='multipart')

    def commit(self, job, index=0, account=None, **flags):
        return self.client.post(self.job_url(job, account), {'batch_index': index, **flags}, format='json')

    def seed(self, price_list, item, price):
        entry = PriceListEntry.objects.filter(price_list=price_list, item=item).first()
        set_prices(self.supplier_a, self.seller_a, changes=[{'list_id': price_list.pk, 'supplier_item_id': item.pk, 'unit_price': Decimal(price),
                                                             'expected_revision': entry.revision if entry else None}])

    def prices(self, price_list):
        return {entry.item.supplier_invent_id: entry.unit_price for entry in PriceListEntry.objects.filter(price_list=price_list).select_related('item')}

    def test_amounts_are_parsed_as_decimal_only_and_ambiguous_or_imprecise_cells_are_rejected(self):
        accepted = {'12,50': '12.50', '$ 12.50': '12.50', '1,234.50': '1234.50', '1.234,50': '1234.50', 'B/. 7': '7.00', 'usd 3.5': '3.50',
                    '12.50 PAB': '12.50', '1,234,567': '1234567.00', '1.234.567,5': '1234567.50', '0': '0.00', 15: '15.00', 15.5: '15.50',
                    # A pasted computed value keeps float noise past Excel's 15 significant digits; it is the amount Excel shows.
                    11.499999999999998: '11.50', 0.30000000000000004: '0.30', 9999999999.99: '9999999999.99', ' $ 12,5 ': '12.50'}
        for value, expected in accepted.items():
            with self.subTest(value=value):
                self.assertEqual(parse_amount(Cell(value, 's' if isinstance(value, str) else 'n')), Decimal(expected))
        for value in [None, '', '   ']:
            self.assertIsNone(parse_amount(Cell(value)))
        self.assertEqual(parse_amount(Cell(' borrar ', 's')), DELETE)
        rejected = {'12,500': 'Formato ambiguo: escribe 12500 o 12.50', '1.234': 'Formato ambiguo: escribe 1234 o 1.23',
                    '12.555': 'Formato ambiguo: escribe 12555 o 12.55', '12.5555': 'Usa un máximo de dos decimales.', '1,234.567': 'Usa un máximo de dos decimales.', 12.555: 'Usa un máximo de dos decimales.',
                    12.345000000000001: 'Usa un máximo de dos decimales.', 0.005: 'Usa un máximo de dos decimales.',
                    '-5': 'El importe no puede ser negativo.', -5: 'El importe no puede ser negativo.', '(5.00)': 'El importe no puede ser negativo.',
                    '$-5': 'El importe no puede ser negativo.', '12345678901': 'Usa como máximo 10 dígitos enteros.', 1e10: 'Usa como máximo 10 dígitos enteros.',
                    'abc': 'Escribe un importe válido, por ejemplo 12.50.', '1 234,50': 'Escribe un importe válido, por ejemplo 12.50.',
                    '12,34,56': 'Escribe un importe válido, por ejemplo 12.50.', float('nan'): 'Escribe un importe válido, por ejemplo 12.50.'}
        for value, message in rejected.items():
            with self.subTest(value=value), self.assertRaisesMessage(ValueError, message):
                parse_amount(Cell(value, 's' if isinstance(value, str) else 'n'))
        for cell, message in [(Cell('=A1*2', 'f'), 'Usa valores, sin fórmulas ni errores de Excel.'), (Cell('#VALUE!', 'e'), 'Usa valores, sin fórmulas'),
                              (Cell(datetime(2026, 1, 2), is_date=True), 'sin fechas ni valores SI/NO'), (Cell(True, 'b'), 'sin fechas ni valores SI/NO')]:
            with self.subTest(cell=cell.value), self.assertRaisesMessage(ValueError, message):
                parse_amount(cell)

    def test_erp_export_with_aliases_and_several_lists_previews_creates_new_lists_once_and_round_trips(self):
        mayorista = PriceList.objects.create(supplier=self.supplier_a, code='MAYORISTA', name='MAYORISTA', created_by=self.seller_a)
        self.seed(self.general, self.item_a, '15.00')
        self.seed(mayorista, self.item_a, '14.00')
        apiag = ingest_inventory(supplier=self.supplier_a, actor=self.seller_a, data={'update_id': 'erp-1', 'supplier_invent_id': 'ERP-9', 'codigo': 'ERP-CODE',
                                                                                       'brand': 'KSM', 'description': 'ERP', 'source': 'apiag', 'quantity': 4})
        headers = ['Supplier Invent ID', 'Code', 'Descripción', 'Familia', 'Precio Piso', 'Precio', 'Precio Mayorista', 'PRECIO_NIVEL_A', 'Costo', 'Ubicación',
                   'Precio Costo', 'ERRORES']
        result = self.upload([['A-001', '58411-1r000-g', 'TAMBOR', 'frenos', '10', '15.00', '13,50', '$ 12.00', 9.87, 'B-12', 9.87, 'viejo'],
                              ['A-002', 'OTRO-CODIGO', 'TAMBOR', '', '', '1.234,50', None, '11', 9.87, '', 9.87, ''],
                              ['ERP-9', 'ERP-CODE', '', '', '', 20, '', '', 9.87, '', 9.87, '']], headers)
        self.assertEqual(result.status_code, 200, result.data)
        data = result.data
        # The documented OpenAPI shape is exactly what the view returns.
        self.assertEqual(set(data), set(PriceImportResult().fields))
        documented = PriceImportResult(data=data)
        self.assertTrue(documented.is_valid(), documented.errors)
        self.assertEqual((data['status'], data['valid'], data['id_column'], data['code_column']), ('ready', True, 'SUPPLIER_INVENT_ID', 'CODE'))
        self.assertEqual(data['ignored_columns'], ['DESCRIPCION', 'COSTO', 'UBICACION', 'PRECIO_COSTO'])
        self.assertEqual(data['columns'], ['SUPPLIER_INVENT_ID', 'CODE', 'DESCRIPCION', 'FAMILIA', 'PRECIO_PISO', 'PRECIO', 'PRECIO_MAYORISTA', 'PRECIO_NIVEL_A'])
        self.assertEqual(data['lists'], [{'header': 'PRECIO', 'code': 'GENERAL', 'new': False, 'archived': False, 'skipped': False},
                                         {'header': 'PRECIO_MAYORISTA', 'code': 'MAYORISTA', 'new': False, 'archived': False, 'skipped': False},
                                         {'header': 'PRECIO_NIVEL_A', 'code': 'NIVEL_A', 'new': True, 'archived': False, 'skipped': False}])
        self.assertEqual(data['lists_to_create'], [{'code': 'NIVEL_A', 'name': 'NIVEL_A', 'currency': 'USD', 'is_default': False, 'header': 'PRECIO_NIVEL_A'}])
        summary = data['summary']
        self.assertEqual({key: summary[key] for key in ['source_rows', 'valid_rows', 'rejected_rows', 'new_prices', 'increases', 'decreases', 'unchanged',
                                                        'removed', 'big_changes', 'line_updates', 'floor_updates']},
                         {'source_rows': 3, 'valid_rows': 3, 'rejected_rows': 0, 'new_prices': 4, 'increases': 0, 'decreases': 1, 'unchanged': 1,
                          'removed': 0, 'big_changes': 0, 'line_updates': 1, 'floor_updates': 1})
        self.assertEqual(summary['lists']['NIVEL_A'], {'new_prices': 2, 'increases': 0, 'decreases': 0, 'unchanged': 0, 'removed': 0})
        self.assertEqual(summary['lists']['MAYORISTA']['decreases'], 1)
        self.assertEqual(data['requires'], {'acknowledge_big_changes': False, 'confirm_new_lists': True})
        self.assertIn({'row': 3, 'column': 'CODE', 'message': 'El CODIGO del archivo (OTRO-CODIGO) no coincide con tu inventario (58411-1R000-KSM); se usó el ID.'},
                      data['warnings'])
        change = next(row for row in data['preview'] if row['list'] == 'MAYORISTA')
        self.assertEqual({key: change[key] for key in ['row', 'supplier_invent_id', 'kind', 'current', 'new', 'change_percent', 'action', 'big_change']},
                         {'row': 2, 'supplier_invent_id': 'A-001', 'kind': 'list_price', 'current': '14.00', 'new': '13.50', 'change_percent': '-3.6',
                          'action': 'update', 'big_change': False})
        self.assertEqual(PriceList.objects.count(), 2, 'Reviewing never writes.')
        refused = self.commit(data['job_id'])
        self.assertEqual((refused.status_code, refused.data), (400, {'detail': 'Confirma la creación de las listas nuevas: NIVEL_A.'}))
        saved = self.commit(data['job_id'], confirm_new_lists=True)
        self.assertEqual((saved.status_code, saved.data['status'], saved.data['imported']), (200, 'completed', True))
        self.assertEqual(saved.data['applied_summary'], {'created': 4, 'updated': 1, 'removed': 0, 'unchanged': 1, 'floor_updates': 1, 'line_updates': 1})
        nivel = PriceList.objects.get(code='NIVEL_A')
        self.assertEqual((nivel.is_default, nivel.currency, nivel.name), (False, 'USD', 'NIVEL_A'))
        self.assertEqual(self.prices(self.general), {'A-001': Decimal('15.00'), 'A-002': Decimal('1234.50'), 'ERP-9': Decimal('20.00')})
        self.assertEqual(self.prices(mayorista), {'A-001': Decimal('13.50')})
        self.assertEqual(self.prices(nivel), {'A-001': Decimal('12.00'), 'A-002': Decimal('11.00')})
        pricing = SupplierItemPricing.objects.get(item=self.item_a)
        self.assertEqual((pricing.discount_group, pricing.floor_price), ('FRENOS', Decimal('10.00')))
        job = data['job_id']
        self.assertEqual(set(PriceChange.objects.filter(source='import').values_list('reference', flat=True)), {f'import:{job}:2', f'import:{job}:3', f'import:{job}:4'})
        self.assertEqual(PriceListEntry.objects.get(price_list=nivel, item=self.item_a).reference, f'import:{job}:2')
        event = PricingAuditEvent.objects.get(kind='price_import_applied')
        self.assertEqual((event.object_id, event.payload['batch'], event.payload['created']), (job, 0, 4))
        self.assertEqual(PricingAuditEvent.objects.filter(kind='price_list_changed', payload__source='import').count(), 1)
        # Supplier cost is never stored anywhere.
        self.assertFalse(PriceChange.objects.filter(new_value=Decimal('9.87')).exists() or PriceListEntry.objects.filter(unit_price=Decimal('9.87')).exists())
        self.assertNotIn('9.87', str(PriceImportJob.objects.values_list('preview', 'summary', 'rejected_rows').get()))
        history = PriceChange.objects.count()
        retry = self.commit(job, confirm_new_lists=True)
        self.assertEqual((retry.status_code, retry.data['progress']['completed_batches']), (200, 1))
        self.assertEqual((PriceChange.objects.count(), PriceList.objects.filter(code='NIVEL_A').count()), (history, 1))
        self.assertEqual(SupplierItem.objects.get(pk=apiag.pk).source, 'apiag')
        # Round trip: the export uploaded as-is changes nothing.
        export = self.client.get(f'/api/v1/accounts/{self.supplier_a.pk}/prices/export/')
        again = self.client.post(self.base(), {'file': SimpleUploadedFile('precios-proveedor.xlsx', export.content)}, format='multipart')
        self.assertEqual(again.status_code, 200, again.data)
        self.assertEqual((again.data['status'], again.data['valid'], again.data['summary']['unchanged'], again.data['preview_count'], again.data['error_count']),
                         ('review', False, 6, 0, 0))
        self.assertEqual(again.data['ignored_columns'], ['MARCA', 'DESCRIPCION'])
        self.assertEqual(again.data['progress']['total_batches'], 0)
        # Edited in Excel and uploaded again: only the edited cell changes.
        book = load_workbook(io.BytesIO(export.content))
        sheet = book['PRECIOS']
        header = [cell.value for cell in sheet[1]]
        row = next(number for number in range(2, sheet.max_row + 1) if sheet.cell(number, 1).value == 'A-002')
        sheet.cell(row, header.index('PRECIO_NIVEL_A') + 1, '11,25')
        output = io.BytesIO()
        book.save(output)
        edited = self.client.post(self.base(), {'file': SimpleUploadedFile('editado.xlsx', output.getvalue())}, format='multipart')
        self.assertEqual((edited.data['preview_count'], edited.data['preview'][0]['new'], edited.data['progress']['total_rows']), (1, '11.25', 1))
        self.assertEqual(self.commit(edited.data['job_id']).status_code, 200)
        self.assertEqual(self.prices(nivel)['A-002'], Decimal('11.25'))

    def test_blank_keeps_borrar_removes_zero_is_skipped_and_rejected_rows_round_trip_through_the_correction_workbook(self):
        self.seed(self.general, self.item_a, '15.00')
        self.seed(self.general, self.item_a2, '8.00')
        set_prices(self.supplier_a, self.seller_a, items=[{'supplier_item_id': self.item_a.pk, 'discount_group': 'FRENOS', 'floor_price': Decimal('10'), 'expected_revision': None}])
        headers = ['ID_INVENTARIO_PROVEEDOR', 'MARCA', 'LINEA', 'PRECIO_MINIMO', 'PRECIO']
        result = self.upload([
            ['A-001', 'KSM', 'borrar', 'BORRAR', None],
            ['A-002', 'KSM', None, None, 0],
            ['NO-EXISTE', 'KSM', '', '', 5.499999999999999],
            ['A-003', 'KSM', '', '', '12,500'],
            ['A-004', 'KSM', '', '', '=1+1'],
            ['A-005', 'KSM', '', '', datetime(2026, 6, 12)],
            ['A-006', 'KSM', '', '', True],
            ['A-007', 'KSM', '', '', -3],
            ['A-008', 'KSM', '', 'x', '4'],
            [datetime(2026, 6, 12), 'KSM', '', '', '4'],
            ['DUP', 'KSM', '', '', '4'],
            ['DUP', 'KSM', '', '', '4'],
        ], headers)
        self.assertEqual(result.status_code, 200, result.data)
        data = result.data
        self.assertEqual((data['summary']['valid_rows'], data['summary']['rejected_rows'], data['summary']['line_updates'], data['summary']['floor_updates']), (2, 10, 1, 1))
        self.assertIn({'row': 3, 'column': 'PRECIO', 'message': 'Precio cero ignorado; deja la celda vacía o escribe BORRAR.'}, data['warnings'])
        messages = {(error['row'], error['column']): error['message'] for error in data['errors']}
        self.assertEqual(messages[(4, 'ID_INVENTARIO_PROVEEDOR')], 'Este ID no existe en tu inventario. Carga primero sus existencias.')
        self.assertEqual(messages[(5, 'PRECIO')], 'Formato ambiguo: escribe 12500 o 12.50')
        self.assertEqual(messages[(6, 'PRECIO')], 'Usa valores, sin fórmulas ni errores de Excel.')
        self.assertEqual(messages[(9, 'PRECIO')], 'El importe no puede ser negativo.')
        self.assertIn('fecha', messages[(11, 'ID_INVENTARIO_PROVEEDOR')])
        self.assertEqual(messages[(12, 'ID_INVENTARIO_PROVEEDOR')], 'El ID DUP está repetido en el archivo. Conserva una sola fila por artículo.')
        self.assertEqual(self.commit(data['job_id']).status_code, 200)
        self.assertEqual(self.prices(self.general), {'A-001': Decimal('15.00'), 'A-002': Decimal('8.00')}, 'Blank and 0 never change a price.')
        pricing = SupplierItemPricing.objects.get(item=self.item_a)
        self.assertEqual((pricing.discount_group, pricing.floor_price, pricing.revision), ('', None, 2))
        removed = self.upload([['A-001', 'BORRAR'], ['A-002', '']], ['ID', 'PRECIO'])
        self.assertEqual((removed.data['summary']['removed'], removed.data['preview'][0]['action']), (1, 'delete'))
        self.assertEqual(self.commit(removed.data['job_id']).status_code, 200)
        self.assertEqual(self.prices(self.general), {'A-002': Decimal('8.00')})
        self.assertEqual(PriceChange.objects.filter(kind='list_price', item=self.item_a).first().new_value, None)
        # The correction workbook keeps every rejected row as text, formulas included, plus an ERRORES column.
        download = self.client.get(f'{self.job_url(data["job_id"])}errors/')
        self.assertEqual((download.status_code, download['Cache-Control']), (200, 'private, no-store'))
        sheet = load_workbook(io.BytesIO(download.content), data_only=False)['PRECIOS']
        self.assertEqual([cell.value for cell in sheet[1]], ['ID_INVENTARIO_PROVEEDOR', 'MARCA', 'LINEA', 'PRECIO_MINIMO', 'PRECIO', 'ERRORES'])
        self.assertEqual(sheet.max_row, 11)
        self.assertEqual((sheet['E4'].value, sheet['E4'].data_type), ('=1+1', 's'))
        self.assertEqual(sheet['E2'].value, '5.50', 'A rejected row keeps the amount that was read, never a float artifact.')
        self.assertIn('Formato ambiguo', sheet['F3'].value)
        for number, (identifier, price) in enumerate([('A-001', '5'), ('A-002', '12500'), ('A-002', '4'), ('A-001', '4'), ('A-001', '4'), ('A-002', '4'),
                                                      ('A-001', '4'), ('NUEVO', '4'), ('DUP-1', '4'), ('DUP-2', '4')], 2):
            sheet.cell(number, 1, identifier)
            sheet.cell(number, 4, '')
            sheet.cell(number, 5, price)
        output = io.BytesIO()
        sheet.parent.save(output)
        repaired = self.client.post(self.base(), {'file': SimpleUploadedFile('CORREGIDO.xlsx', output.getvalue())}, format='multipart')
        self.assertEqual(repaired.status_code, 200, repaired.data)
        self.assertEqual(repaired.data['ignored_columns'], ['MARCA'], 'ERRORES is skipped silently.')
        self.assertEqual(repaired.data['summary']['rejected_rows'], 10, 'Repeated IDs are rejected in the corrected file too.')

    def test_files_without_usable_headers_are_rejected_before_any_job_exists(self):
        cases = [(['ID_INVENTARIO_PROVEEDOR', 'CODIGO', 'MARCA', 'COSTO'], 'No encontramos columnas de precio. Usa PRECIO, PRECIO_<LISTA>, LINEA o PRECIO_MINIMO.'),
                 (['CODIGO', 'PRECIO'], 'Se requiere ID_INVENTARIO_PROVEEDOR en la primera fila.'),
                 (['ID', 'PRECIO', 'PRECIO_GENERAL'], 'Las columnas PRECIO y PRECIO_GENERAL corresponden a la misma lista GENERAL. Deja solo una.'),
                 (['ID', 'ID_INVENTARIO_PROVEEDOR', 'PRECIO'], 'La columna ID_INVENTARIO_PROVEEDOR está repetida. Usa cada encabezado una sola vez.'),
                 (['ID', 'LINEA', 'FAMILIA'], 'La columna FAMILIA está repetida. Usa cada encabezado una sola vez.'),
                 (['ID', 'PRECIO_NIVEL-A'], 'El encabezado PRECIO_NIVEL-A no corresponde a una lista'),
                 (['ID', 'PRECIO_'], 'El encabezado PRECIO_ no corresponde a una lista'),
                 (['ID', 'PRECIO'] + [f'EXTRA_{index}' for index in range(39)], 'Usa como máximo 40 columnas en el archivo de precios.')]
        for headers, message in cases:
            with self.subTest(headers=headers[:3]):
                response = self.upload([['A-001'] + ['1'] * (len(headers) - 1)], headers)
                self.assertEqual(response.status_code, 400)
                self.assertIn(message, str(response.data))
        self.assertEqual(self.client.post(self.base(), {'file': SimpleUploadedFile('broken.xlsx', b'broken')}, format='multipart').status_code, 400)
        self.assertEqual(self.client.post(self.base(), {'file': self.file([['A-001', '1', '2']], name='precios.csv')}, format='multipart').status_code, 400)
        self.assertEqual(self.upload([], ['ID', 'PRECIO']).status_code, 400)
        # The PRECIOS sheet wins over the first sheet; LISTA_PRECIOS is accepted too.
        def other_first(sheet):
            sheet.parent.create_sheet('NOTAS', 0)
        listed = self.upload([['A-001', '9']], ['ID', 'PRECIO'], title='Lista Precios', customize=other_first)
        self.assertEqual((listed.status_code, listed.data['summary']['new_prices']), (200, 1))
        self.assertEqual(PriceImportJob.objects.count(), 1)

    def test_suppliers_without_lists_get_a_default_list_from_precio_and_numeric_ids_keep_their_zeros(self):
        PriceList.objects.all().delete()
        SupplierPricingSettings.objects.create(supplier=self.supplier_a, default_currency='PAB')
        SupplierItem.objects.filter(pk=self.item_a2.pk).update(supplier_invent_id='00045')
        result = self.upload([[45, '4,50', '5'], ['A-001', '', '6']], ['ID', 'PRECIO_TALLER', 'PRECIO'],
                             customize=lambda sheet: setattr(sheet['A2'], 'number_format', '00000'))
        self.assertEqual(result.status_code, 200, result.data)
        self.assertEqual([(spec['code'], spec['currency'], spec['is_default']) for spec in result.data['lists_to_create']],
                         [('TALLER', 'PAB', False), ('GENERAL', 'PAB', True)])
        self.assertIn({'row': 2, 'column': 'ID', 'message': 'Identificador numérico convertido a texto: 00045.'}, result.data['warnings'])
        self.assertEqual(self.commit(result.data['job_id'], confirm_new_lists=True).status_code, 200)
        self.assertEqual(list(PriceList.objects.order_by('code').values_list('code', 'is_default', 'currency')), [('GENERAL', True, 'PAB'), ('TALLER', False, 'PAB')])
        self.assertEqual(self.prices(PriceList.objects.get(code='GENERAL')), {'00045': Decimal('5.00'), 'A-001': Decimal('6.00')})

    def test_archived_lists_warn_and_columns_without_new_prices_create_no_list(self):
        PriceList.objects.create(supplier=self.supplier_a, code='VIEJA', name='VIEJA', active=False, created_by=self.seller_a)
        result = self.upload([['A-001', '5', '', 'BORRAR']], ['ID', 'PRECIO_VIEJA', 'PRECIO_VACIA', 'PRECIO_SOLO_BORRAR'])
        self.assertEqual(result.status_code, 200, result.data)
        self.assertEqual([(row['code'], row['new'], row['archived'], row['skipped']) for row in result.data['lists']],
                         [('VIEJA', False, True, False), ('VACIA', False, False, True), ('SOLO_BORRAR', False, False, True)])
        self.assertEqual((result.data['lists_to_create'], result.data['requires']['confirm_new_lists']), ([], False))
        self.assertIn('La lista VIEJA está archivada: sus precios se guardan, pero no se sugieren en cotizaciones.', [warning['message'] for warning in result.data['warnings']])
        self.assertEqual(self.commit(result.data['job_id']).status_code, 200)
        self.assertEqual(self.prices(PriceList.objects.get(code='VIEJA')), {'A-001': Decimal('5.00')})
        self.assertEqual(sorted(PriceList.objects.values_list('code', flat=True)), ['GENERAL', 'VIEJA'])

    def test_an_existing_list_coded_like_a_cost_column_round_trips(self):
        costo = PriceList.objects.create(supplier=self.supplier_a, code='COSTO', name='COSTO', created_by=self.seller_a)
        self.seed(costo, self.item_a, '5.00')
        template = load_workbook(io.BytesIO(self.client.get(f'{self.base()}template/').content))['PRECIOS']
        self.assertIn('PRECIO_COSTO', [cell.value for cell in template[1]])
        result = self.upload([['A-001', '6.00', '7.00']], ['ID', 'PRECIO_COSTO', 'PRECIO_COMPRA'])
        self.assertEqual((result.status_code, result.data['ignored_columns'], result.data['summary']['lists']['COSTO']['increases']),
                         (200, ['PRECIO_COMPRA'], 1), 'Only a cost column without a list of that code is ignored.')
        self.assertEqual(self.commit(result.data['job_id']).status_code, 200)
        self.assertEqual(self.prices(costo), {'A-001': Decimal('6.00')})

    def test_big_changes_require_an_explicit_acknowledgement(self):
        self.seed(self.general, self.item_a, '10.00')
        self.seed(self.general, self.item_a2, '10.00')
        result = self.upload([['A-001', '', '16.00'], ['A-002', '', '15']])
        self.assertEqual((result.data['summary']['big_changes'], result.data['summary']['increases']), (1, 2), 'Exactly 50 % is not a big change.')
        self.assertEqual(result.data['preview'][0], {'row': 2, 'supplier_invent_id': 'A-001', 'codigo': '58411-1R000-G', 'description': 'TAMBOR A-001',
                                                     'kind': 'list_price', 'list': 'GENERAL', 'current': '10.00', 'new': '16.00', 'change_percent': '60.0',
                                                     'action': 'update', 'big_change': True})
        self.assertIn('Variación mayor al 50 % (+60.0 %): 10.00 → 16.00.', [warning['message'] for warning in result.data['warnings']])
        self.assertTrue(result.data['requires']['acknowledge_big_changes'])
        refused = self.commit(result.data['job_id'])
        self.assertEqual((refused.status_code, refused.data), (400, {'detail': 'Confirma que revisaste los cambios mayores al 50 %.'}))
        self.assertEqual(self.prices(self.general)['A-001'], Decimal('10.00'))
        saved = self.commit(result.data['job_id'], acknowledge_big_changes=True)
        self.assertEqual((saved.status_code, saved.data['requires']['acknowledge_big_changes']), (200, False))
        self.assertEqual(self.prices(self.general), {'A-001': Decimal('16.00'), 'A-002': Decimal('15.00')})

    def test_prices_changed_after_preview_return_409_while_stock_syncs_never_make_a_batch_stale(self):
        self.seed(self.general, self.item_a, '10.00')
        before = (list(SupplierItem.objects.order_by('pk').values()), StockEntry.objects.count(), InventoryUpdate.objects.count(), list(MatchingQueue.objects.values()))
        stock_rows = [{'row': 2, 'data': {'supplier_invent_id': 'A-001', 'codigo': '58411-1R000-G', 'brand': 'KSM', 'quantity': 4}}]
        stock_fingerprint = stock_plan(self.supplier_a, [dict(row, data=dict(row['data'])) for row in stock_rows])[1]
        for mutate in [lambda: self.seed(self.general, self.item_a, '10.50'),
                       lambda: set_prices(self.supplier_a, self.seller_a, items=[{'supplier_item_id': self.item_a.pk, 'floor_price': Decimal('1'), 'expected_revision': None}]),
                       lambda: self.seed(self.general, self.item_a2, '3.00')]:
            result = self.upload([['A-001', '', '11.00'], ['A-002', '', '4.00']], ['ID', 'CODIGO', 'PRECIO'])
            mutate()
            history = PriceChange.objects.count()
            response = self.commit(result.data['job_id'])
            self.assertEqual((response.status_code, response.data['detail']), (409, 'Los precios cambiaron desde la revisión; vuelve a subir el archivo.'))
            self.assertEqual(PriceChange.objects.count(), history)
            self.assertEqual(PriceImportJob.objects.get(pk=result.data['job_id']).completed_batches, 0)
        self.assertEqual(stock_plan(self.supplier_a, [dict(row, data=dict(row['data'])) for row in stock_rows])[1], stock_fingerprint)
        self.assertEqual((list(SupplierItem.objects.order_by('pk').values()), StockEntry.objects.count(), InventoryUpdate.objects.count(),
                          list(MatchingQueue.objects.values())), before, 'Price writes never touch stock rows, the ledger or matching.')
        result = self.upload([['A-001', '', '11.00']], ['ID', 'CODIGO', 'PRECIO'])
        ingest_inventory(supplier=self.supplier_a, actor=self.seller_a, data={'update_id': 'stock-sync', 'supplier_invent_id': 'A-001', 'codigo': '58411-1R000-G',
                                                                              'brand': 'KSM', 'description': 'TAMBOR NUEVO', 'source': 'upload', 'quantity': 1})
        self.assertEqual(self.commit(result.data['job_id']).status_code, 200)
        self.assertEqual(self.prices(self.general)['A-001'], Decimal('11.00'))

    def test_batches_checkpoint_committed_indexes_and_a_failed_first_batch_never_leaves_a_new_list(self):
        SupplierItem.objects.bulk_create(SupplierItem(supplier=self.supplier_a, supplier_invent_id=f'BULK-{index:04}', codigo=f'C-{index}', brand='KSM',
                                                      source='upload') for index in range(1001))
        result = self.upload([[f'BULK-{index:04}', '', f'{index + 1}.25'] for index in range(1001)], ['ID', 'CODIGO', 'PRECIO_NUEVA'])
        self.assertEqual(result.status_code, 200)
        job = result.data['job_id']
        self.assertEqual((result.data['progress']['total_batches'], result.data['progress']['batch_size'], result.data['preview_count'], len(result.data['preview'])),
                         (2, 1000, 1001, 100))
        self.assertEqual(PriceImportBatch.objects.filter(job_id=job).count(), 2)
        with patch('mall.price_import.set_prices', side_effect=ValidationError('Conflicto simulado.')):
            self.assertEqual(self.commit(job, confirm_new_lists=True).status_code, 400)
        self.assertFalse(PriceList.objects.filter(code='NUEVA').exists())
        self.assertEqual(PriceImportJob.objects.get(pk=job).completed_batches, 0)
        self.assertEqual(self.commit(job, 1, confirm_new_lists=True).status_code, 409)
        first = self.commit(job, confirm_new_lists=True)
        self.assertEqual((first.status_code, first.data['status'], first.data['progress']['processed_rows'], first.data['progress']['next_batch']),
                         (200, 'importing', 1000, 1))
        self.assertEqual(self.commit(job, confirm_new_lists=True).data['progress']['processed_rows'], 1000)
        self.assertEqual(self.client.get(self.job_url(job)).data['requires'], {'acknowledge_big_changes': False, 'confirm_new_lists': False})
        self.assertEqual(self.commit(job, 2).status_code, 409)
        last = self.commit(job, 1)
        self.assertEqual((last.status_code, last.data['status'], last.data['applied_summary']['created']), (200, 'completed', 1001))
        self.assertEqual(PriceList.objects.filter(code='NUEVA').count(), 1)
        self.assertEqual(PriceListEntry.objects.filter(price_list__code='NUEVA').count(), 1001)
        self.assertEqual(PricingAuditEvent.objects.filter(kind='price_import_applied').count(), 2)

    def test_jobs_are_owner_bound_permissioned_and_expire_but_keep_their_checkpoint_and_errors(self):
        result = self.upload([['A-001', '', '9.00'], ['NO-EXISTE', '', '1']])
        job = result.data['job_id']
        colleague = request_tests.User.objects.create_user('price-colleague', 'price-colleague@example.invalid', 'Request938!')
        Membership.objects.create(user=colleague, account=self.supplier_a)
        self.client.force_authenticate(colleague)
        for response in [self.client.get(self.job_url(job)), self.commit(job), self.client.get(f'{self.job_url(job)}errors/')]:
            self.assertEqual(response.status_code, 404)
        for user, account, status in [(self.seller_b, self.supplier_a, 404), (self.root, self.supplier_a, 404), (self.buyer, self.client_account, 403),
                                      (self.seller_a, self.supplier_b, 404)]:
            self.client.force_authenticate(user)
            with self.subTest(user=user.username):
                self.assertEqual(self.client.get(self.job_url(job, account)).status_code, status)
                self.assertEqual(self.commit(job, account=account).status_code, status)
                self.assertEqual(self.upload([['A-001', '', '1']], account=account).status_code, status)
                self.assertEqual(self.client.get(f'{self.base(account)}template/').status_code, status)
        self.client.force_authenticate(None)
        self.assertEqual(self.client.get(f'{self.base()}template/').status_code, 401)
        self.assertEqual(self.commit(job).status_code, 401)
        # Staff below config_min_permission read the template but cannot upload, read or apply imports.
        SupplierPricingSettings.objects.create(supplier=self.supplier_a, config_min_permission='manager')
        self.client.force_authenticate(self.seller_a)
        self.assertEqual(self.client.get(f'{self.base()}template/').status_code, 200)
        for response in [self.upload([['A-001', '', '1']]), self.client.get(self.job_url(job)), self.commit(job), self.client.get(f'{self.job_url(job)}errors/')]:
            self.assertEqual(response.status_code, 403)
        Membership.objects.filter(user=self.seller_a, account=self.supplier_a).update(permission='manager')
        PriceImportJob.objects.filter(pk=job).update(expires_at=timezone.now() - timedelta(days=1))
        expired = self.commit(job)
        self.assertEqual((expired.status_code, expired.data['detail']), (409, 'Esta revisión venció. Sus errores siguen disponibles para descargar; vuelve a subir el archivo.'))
        self.assertEqual(self.client.get(self.job_url(job)).status_code, 200)
        self.assertEqual(self.client.get(f'{self.job_url(job)}errors/').status_code, 200)
        PriceImportJob.objects.filter(pk=job).update(expires_at=timezone.now() + timedelta(days=1))
        self.assertEqual(self.commit(job).status_code, 200)
        PriceImportJob.objects.filter(pk=job).update(expires_at=timezone.now() - timedelta(days=1))
        self.assertEqual(self.commit(job).status_code, 200, 'A committed batch replays its checkpoint even after expiry.')
        self.assertEqual(self.prices(self.general), {'A-001': Decimal('9.00')})

    def test_template_has_one_sample_row_text_ids_and_a_column_per_active_list(self):
        PriceList.objects.create(supplier=self.supplier_a, code='MAYORISTA', name='MAYORISTA', created_by=self.seller_a)
        PriceList.objects.create(supplier=self.supplier_a, code='VIEJA', name='VIEJA', active=False, created_by=self.seller_a)
        response = self.client.get(f'{self.base()}template/')
        self.assertEqual((response.status_code, response['Content-Type'], response['Cache-Control']), (200, XLSX, 'private, no-store'))
        sheet = load_workbook(io.BytesIO(response.content))['PRECIOS']
        rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
        self.assertEqual(rows[0], ['ID_INVENTARIO_PROVEEDOR', 'CODIGO', 'MARCA', 'DESCRIPCION', 'LINEA', 'PRECIO_MINIMO', 'PRECIO', 'PRECIO_MAYORISTA'])
        self.assertEqual(rows[1:], [['000123', '58411-1R000-KSM', 'KSM', 'REEMPLAZA ESTA FILA CON TUS DATOS', 'FRENOS', 12.5, 15.75, None]])
        self.assertEqual((sheet['A2'].number_format, sheet['A2'].data_type), ('@', 's'))
        unchanged = self.client.post(self.base(), {'file': SimpleUploadedFile('plantilla.xlsx', response.content)}, format='multipart')
        self.assertEqual((unchanged.status_code, unchanged.data['errors'][0]['message']), (200, 'Este ID no existe en tu inventario. Carga primero sus existencias.'))

    def test_a_running_price_import_never_marks_the_supplier_busy_for_matching(self):
        PriceImportJob.objects.create(owner=self.seller_a, supplier=self.supplier_a, filename='precios.xlsx', status='importing',
                                      expires_at=timezone.now() + timedelta(days=1), total_batches=2)
        description = 'TAMBOR KIA PICANTO 04-11'
        def stock(stable, code, brand=''):
            return ingest_inventory(supplier=self.supplier_a, actor=self.seller_a, data={'update_id': stable, 'supplier_invent_id': stable, 'codigo': code,
                                                                                         'brand': brand, 'description': description, 'source': 'upload', 'quantity': 1})
        pending, reviewed, blocked = stock('PEND-1', '58411-07000-K'), stock('REV-1', 'D-HYU-1R', 'ACME'), stock('REV-2', 'D-KIA-2R', 'ACME')
        reconcile_items(MatchIndex())
        target = Part.objects.create(sku='58411-07000', description=description)
        # The worker's supplier pass (reconcile_items) still examines this supplier while its prices import.
        self.assertIsNotNone(process_queue(force=True, use_ai=False))
        pending.refresh_from_db()
        self.assertEqual((pending.part_id, pending.matching_status), (target.pk, 'matched'))
        apply_item(MatchingCase.objects.get(item=reviewed), part_row(target))
        reviewed.refresh_from_db()
        self.assertEqual((reviewed.part_id, reviewed.matching_status), (target.pk, 'matched'))
        # A stock import of the same supplier does pause matching: the guard is real, a price import just never trips it.
        SupplierInventoryImportJob.objects.create(owner=self.seller_a, supplier=self.supplier_a, filename='stock.xlsx', status='importing',
                                                  expires_at=timezone.now() + timedelta(days=1))
        with self.assertRaisesMessage(ValidationError, 'Espera a que termine la importación de este proveedor.'):
            apply_item(MatchingCase.objects.get(item=blocked), part_row(target))


@skipUnless(connection.vendor == 'postgresql', 'Row locks are only meaningful on PostgreSQL.')
class PriceImportConcurrencyTests(APITransactionTestCase):
    setUp = deal_tests.ConcurrentDealTests.setUp
    call = deal_tests.ConcurrentDealTests.call

    def test_an_import_batch_and_a_manual_edit_of_the_same_price_give_one_winner(self):
        general = PriceList.objects.create(supplier=self.supplier, code='GENERAL', name='GENERAL', is_default=True, created_by=self.seller)
        set_prices(self.supplier, self.seller, changes=[{'list_id': general.pk, 'supplier_item_id': self.item.pk, 'unit_price': Decimal('10.00'), 'expected_revision': None}])
        book = Workbook()
        book.active.append(['ID', 'PRECIO'])
        book.active.append(['1', '12.00'])
        output = io.BytesIO()
        book.save(output)
        self.client.force_authenticate(self.seller)
        base = f'/api/v1/accounts/{self.supplier.pk}'
        job = self.client.post(f'{base}/prices/import/', {'file': SimpleUploadedFile('p.xlsx', output.getvalue())}, format='multipart').data['job_id']
        calls = [(f'{base}/prices/import/jobs/{job}/', {'batch_index': 0}),
                 (f'{base}/prices/', {'changes': [{'list_id': str(general.pk), 'supplier_item_id': str(self.item.pk), 'unit_price': '11.00', 'expected_revision': 1}]})]
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda call: self.call(self.seller, *call)[0], calls))
        self.assertEqual(sorted(results), [200, 409])
        entry = PriceListEntry.objects.get()
        self.assertEqual((entry.revision, entry.unit_price), (2, Decimal('12.00') if results[0] == 200 else Decimal('11.00')))
        self.assertEqual(PriceChange.objects.count(), 2)
