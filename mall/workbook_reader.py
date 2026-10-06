"""Shared, defensive .xlsx reading and writing for supplier imports (stock and prices).

Uploads are read with read_only/data_only=False/keep_links=False behind a 30 MB uncompressed-size guard, so formulas stay
visible (and are rejected) instead of being evaluated, and external links are never followed. Correction workbooks write
every user-supplied cell as text, so a rejected formula can never execute when the file is opened again.
"""
import io
import math
import re
import unicodedata
from contextlib import contextmanager
from decimal import Decimal
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from rest_framework.exceptions import ValidationError

XLSX_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
MAX_UNCOMPRESSED = 30 * 1024 * 1024
READ_ERRORS = (ValueError, KeyError, OSError, IndexError, ParseError, DefusedXmlException, BadZipFile)


def header(value):
    text = ''.join(c for c in unicodedata.normalize('NFD', str(value or '').strip().upper()) if not unicodedata.combining(c))
    return '_'.join(text.split())


def identifier_text(cell):
    """Read numeric codes without rounding or applying non-identifier formats."""
    value = cell.value
    if isinstance(value, str):
        return value
    if cell.is_date:
        raise ValueError('Excel convirtió este identificador en una fecha. Recupera el código original y guárdalo como texto.')
    if isinstance(value, bool) or not isinstance(value, (int, float)) or (isinstance(value, float) and not math.isfinite(value)):
        raise ValueError('Usa un código de texto o un número válido, sin fechas ni valores SI/NO.')
    number = Decimal(str(value))
    if abs(number) >= Decimal('1e15') or len(number.normalize().as_tuple().digits) > 15:
        raise ValueError('Este código numérico puede haber perdido precisión en Excel. Recupera el código original y guárdalo como texto.')
    plain = format(number, 'f')
    integer, dot, fraction = plain.partition('.')
    fraction = fraction.rstrip('0')
    plain = integer + ('.' + fraction if fraction else '')
    mask = cell.number_format or 'General'
    if mask.lower() in ['general', '@']:
        return plain
    match = re.fullmatch(r'(0+)(?:\.(0+#*|#+))?', mask)
    if not match:
        raise ValueError('El formato numérico de este código no es compatible. Guarda el código original como texto.')
    decimal_mask = match.group(2) or ''
    if len(fraction) > len(decimal_mask):
        raise ValueError('El formato de Excel redondea este código. Recupera su valor original y guárdalo como texto.')
    integer = integer.zfill(len(match.group(1)) + (1 if integer.startswith('-') else 0))
    fraction = fraction.ljust(decimal_mask.count('0'), '0')
    return integer + ('.' + fraction if fraction else '')


def open_workbook(raw):
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED:
                raise ValidationError('El contenido del archivo Excel es demasiado grande.')
        return load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, ValueError, KeyError, OSError, ParseError, DefusedXmlException, InvalidFileException):
        raise ValidationError('No pudimos leer el archivo. Usa un libro Excel .xlsx válido.')


@contextmanager
def workbook_rows(raw, *, sheets, max_columns, too_wide):
    """Rows of the first sheet named in `sheets` (or the first sheet). Streams actual rows: some exporters report incomplete
    dimensions, and no populated row may disappear because of that metadata. Any low-level read error inside the block becomes
    one Spanish validation error."""
    workbook = open_workbook(raw)
    try:
        sheet = next((sheet for sheet in workbook.worksheets if header(sheet.title) in sheets), workbook.worksheets[0])
        if (sheet.max_column or 0) > max_columns:
            raise ValidationError(too_wide)
        sheet.reset_dimensions()
        yield sheet.iter_rows(max_col=max_columns)
    except READ_ERRORS:
        raise ValidationError('No pudimos leer las filas del archivo Excel.')
    finally:
        workbook.close()


def cell_value(cell):
    value = cell.value
    return value.isoformat() if hasattr(value, 'isoformat') else str(value) if value is not None else ''


def cell_original(cell):
    """What the uploaded cell really held, kept with rejected rows for their correction workbook."""
    return {'value': cell_value(cell), 'number_format': cell.number_format,
            'kind': 'formula' if cell.data_type == 'f' else 'error' if cell.data_type == 'e' else
            'date' if cell.is_date else 'boolean' if isinstance(cell.value, bool) else
            'text' if isinstance(cell.value, str) else 'number' if cell.value is not None else 'blank'}


def reject_formula(cell):
    if cell.data_type in ['f', 'e']:
        raise ValueError('Usa valores, sin fórmulas ni errores de Excel.')


def reject_date_or_boolean(cell, message='Usa texto, sin fechas ni valores SI/NO.'):
    if cell.is_date or isinstance(cell.value, bool):
        raise ValueError(message)


def reject_repeated(records, field, column, message):
    """Reject every occurrence of a repeated stable ID: never add, merge or choose one row arbitrarily, even when two copies look
    identical. `message` is formatted with {identifier}."""
    seen = {}
    for record in records:
        identifier = record['data'].get(field)
        if identifier:
            seen.setdefault(identifier, []).append(record)
    for identifier, occurrences in seen.items():
        if len(occurrences) > 1:
            for record in occurrences:
                record['errors'].append({'row': record['row'], 'column': column, 'message': message.format(identifier=identifier)})


def xlsx_response(book, filename):
    output = io.BytesIO()
    book.save(output)
    response = HttpResponse(output.getvalue(), content_type=XLSX_TYPE)
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response['Cache-Control'] = 'private, no-store'
    return response


def workbook_response(title, columns, rows, filename, *, widths, numeric=(), render=None, errors=False):
    """A template or correction workbook. columns: [(field, HEADER)]; rows: dicts, or {'values': dict, 'errors': [...]}.
    Every cell is explicit text except `numeric` fields of a template: text keeps initial zeros and prevents a rejected formula
    from executing when its repair workbook is opened."""
    book = Workbook()
    sheet = book.active
    sheet.title = title
    sheet.append([name for _, name in columns] + (['ERRORES'] if errors else []))
    for number, row in enumerate(rows, 2):
        values = row.get('values', row)
        for column, (field, _) in enumerate(columns, 1):
            cell = sheet.cell(number, column)
            value = values.get(field, '')
            cell.value = render(field, value) if render else value
            if field not in numeric or errors:
                cell.data_type = 's'
                cell.number_format = '@'
        if errors:
            cell = sheet.cell(number, len(columns) + 1, '; '.join(f'{error["column"]}: {error["message"]}' for error in row['errors']))
            cell.data_type = 's'
    sheet.freeze_panes = 'A2'
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    return xlsx_response(book, filename)


def job_progress(job, batch_size):
    return {'batch_size': batch_size, 'total_batches': job.total_batches, 'completed_batches': job.completed_batches,
            'total_rows': job.total_rows, 'processed_rows': job.processed_rows,
            'next_batch': job.completed_batches if job.completed_batches < job.total_batches else None}


def batches(rows, size):
    return [(offset // size, rows[offset:offset + size]) for offset in range(0, len(rows), size)]
