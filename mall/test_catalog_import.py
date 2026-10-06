import io
from datetime import datetime
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError
from openpyxl import Workbook
from rest_framework.test import APITestCase

from .models import Account, CatalogImportJob, Membership, Part, PartCode, Role, StockEntry, SupplierItem, User
from .services import ingest_inventory


class CatalogImportTests(APITestCase):
    url = '/api/v1/management/catalog/import/'
    columns = ['SKU', 'NOMBRE', 'DESCRIPCION', 'CODIGO_ALTERNO', 'MARCA_ALTERNO', 'ACTIVO']

    def setUp(self):
        self.root = User.objects.create_superuser('excel-root', 'excel-root@example.invalid', 'ExcelRoot938!')
        self.user = User.objects.create_user('excel-member', 'excel-member@example.invalid', 'ExcelMember938!')
        self.client.force_authenticate(self.root)

    def excel(self, rows, headers=None):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Inventario'
        sheet.append(headers or self.columns)
        for row in rows:
            sheet.append(row)
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def upload(self, raw, **options):
        return self.client.post(self.url, {'file': SimpleUploadedFile('inventario.xlsx', raw), **options}, format='multipart')

    def test_only_current_superusers_can_preview_or_import(self):
        raw = self.excel([['SKU-1']])
        self.client.force_authenticate(None)
        self.assertEqual(self.upload(raw).status_code, 401)
        self.user.is_staff = True
        self.user.save()
        self.client.force_authenticate(self.user)
        for mode in ['preview', 'import']:
            self.assertEqual(self.upload(raw, mode=mode).status_code, 403)
        self.assertEqual(Part.objects.count(), 0)

    def test_preview_then_import_groups_alternos_preserves_ids_codes_and_ledger(self):
        existing = Part.objects.create(sku='58411-1R000', name='NOMBRE EXISTENTE', description='DESCRIPCIÓN EXISTENTE', active=True)
        old_code = PartCode.objects.create(part=existing, code='OLD-CODE')
        supplier = Account.objects.create(name='Proveedor Excel')
        supplier.roles.add(Role.objects.get(code='supplier_retail'))
        Membership.objects.create(user=self.user, account=supplier)
        stock = ingest_inventory(supplier=supplier, actor=self.user, data={'update_id': 'before-import', 'supplier_invent_id': 'stable', 'codigo': 'OLD-CODE', 'brand': 'OEM', 'source': 'upload', 'quantity': 7})
        ledger_before = list(StockEntry.objects.values())
        raw = self.excel([
            ['58411-1r000', '', '', '58411-1r000-g'],
            ['58411-1R000', '', '', '58-1r0'],
            ['58411-1R000', '', '', '58-1r0'],
            ['000123', 'repuesto nuevo', 'descripción', '000045', 'oem', 'NO'],
        ])
        preview = self.upload(raw)
        self.assertEqual(preview.status_code, 200)
        self.assertTrue(preview.data['valid'])
        self.assertEqual(preview.data['summary'], {'rows': 4, 'created_skus': 1, 'updated_skus': 1, 'unchanged_skus': 0, 'added_codes': 3})
        self.assertEqual(Part.objects.count(), 1)
        self.assertEqual(PartCode.objects.count(), 1)
        imported = self.upload(raw, mode='import', preview_token=preview.data['preview_token'])
        self.assertEqual(imported.status_code, 200)
        self.assertTrue(imported.data['imported'])
        existing.refresh_from_db()
        self.assertEqual(existing.name, 'NOMBRE EXISTENTE')
        self.assertEqual(existing.description, 'DESCRIPCIÓN EXISTENTE')
        self.assertTrue(existing.active)
        self.assertTrue(PartCode.objects.filter(pk=old_code.pk).exists())
        self.assertEqual(existing.codes.count(), 3)
        new = Part.objects.get(sku='000123')
        self.assertEqual(new.name, 'REPUESTO NUEVO')
        self.assertFalse(new.active)
        self.assertEqual(new.codes.get().code, '000045')
        stock.refresh_from_db()
        self.assertEqual(stock.part_id, existing.pk)
        self.assertEqual(stock.reported_quantity, 7)
        self.assertEqual(stock.supplier_invent_id, 'stable')
        self.assertEqual(list(StockEntry.objects.values()), ledger_before)
        repeated = self.upload(raw)
        self.assertEqual(repeated.data['summary']['unchanged_skus'], 2)
        self.assertEqual(repeated.data['summary']['added_codes'], 0)
        self.assertEqual(self.upload(raw, mode='import', preview_token=repeated.data['preview_token']).status_code, 200)
        self.assertEqual(PartCode.objects.count(), 4)

    def test_catalog_and_in_file_code_conflicts_report_exact_rows_without_writes(self):
        first = Part.objects.create(sku='FIRST')
        PartCode.objects.create(part=first, brand='OEM', code='SHARED')
        raw = self.excel([['SECOND', '', '', 'SHARED'], ['THIRD', '', '', 'FIRST'], ['FOURTH', '', '', 'FIFTH'], ['FIFTH']])
        result = self.upload(raw)
        self.assertEqual(result.status_code, 200)
        self.assertFalse(result.data['valid'])
        self.assertEqual({error['row'] for error in result.data['errors']}, {2, 3, 4, 5})
        self.assertNotIn('preview_token', result.data)
        self.assertEqual(Part.objects.count(), 1)
        self.assertEqual(PartCode.objects.count(), 1)

    def test_repeated_sku_metadata_and_invalid_identifiers_are_rejected(self):
        raw = self.excel([['SKU-A', 'FIRST'], ['SKU-A', 'SECOND'], [datetime(2026, 6, 12)], ['VALID', '', '', '=A2'], ['BRAND-ONLY', '', '', '', 'OEM'], ['BAD-ACTIVE', '', '', '', '', 'MAYBE']])
        result = self.upload(raw)
        self.assertFalse(result.data['valid'])
        self.assertEqual({error['row'] for error in result.data['errors']}, {2, 3, 4, 5, 6, 7})
        self.assertEqual(next(error for error in result.data['errors'] if error['row'] == 3)['column'], 'NOMBRE')
        self.assertEqual(result.data['status'], 'review')
        self.assertEqual(result.data['pending_error_count'], 6)
        self.assertEqual(Part.objects.count(), 0)

    def test_numeric_identifiers_preserve_text_zero_padding_and_decimal_values(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Inventario'
        sheet.append(self.columns)
        sheet.append([146690, '', '', 45])
        sheet['D2'].number_format = '000000'
        sheet.append([5035.27])
        sheet.append([7075133.5])
        sheet.append([12.34])
        sheet['A5'].number_format = '000.0000'
        sheet.append(['000123', '', '', 'TEXT-000045'])
        sheet.append([999999999999999])
        sheet.append([1e-7])
        stream = io.BytesIO()
        workbook.save(stream)
        raw = stream.getvalue()
        preview = self.upload(raw).data
        self.assertTrue(preview['valid'])
        self.assertEqual(preview['warning_count'], 7)
        self.assertEqual([item['sku'] for item in preview['preview']], ['146690', '5035.27', '7075133.5', '012.3400', '000123', '999999999999999', '0.0000001'])
        self.assertEqual(preview['preview'][0]['codes'], [{'code': '000045', 'brand': ''}])
        job_url = f"{self.url}jobs/{preview['job_id']}/"
        saved = self.client.post(job_url, {'batch_index': 0})
        self.assertTrue(saved.data['imported'])
        self.assertEqual(saved.data['warning_count'], 7)
        self.assertEqual(Part.objects.get(sku='146690').codes.get().code, '000045')
        self.assertTrue(Part.objects.filter(sku='000123').exists())
        repeated = self.upload(raw).data
        self.assertEqual(repeated['summary']['unchanged_skus'], 7)

    def test_numeric_identifier_formats_never_silently_round_dates_or_precision(self):
        from .catalog_import import identifier_text
        sheet = Workbook().active
        supported = [(123, 'General', '123'), (123, '@', '123'), (123, '0', '123'),
                     (123, '000000', '000123'), (12.34, '0.0000', '12.3400'),
                     (-12, '0000', '-0012'), (12.3, '0.00##', '12.30'),
                     ('000123', 'General', '000123'), ('1E10', 'General', '1E10'),
                     ('12345678901234567890', 'General', '12345678901234567890')]
        for value, mask, expected in supported:
            cell = sheet['A1']; cell.value = value; cell.number_format = mask
            self.assertEqual(identifier_text(cell), expected)
        unsupported = [(datetime(2026, 6, 12), 'd-mmm'), (True, 'General'),
                       (1000000000000000, 'General'), (12.345, '0.00'), (12.3, '0'),
                       (12, '0%'), (12, '$0.00'), (12, '0.00E+00'), (12, '# ?/?'),
                       (12, '0;[Red]-0'), (0.30000000000000004, 'General')]
        for value, mask in unsupported:
            cell = sheet['A1']; cell.value = value; cell.number_format = mask
            with self.assertRaises(ValueError, msg=f'{value}: {mask}'):
                identifier_text(cell)

    def test_numeric_conversion_warnings_remain_visible_when_date_codes_are_deferred(self):
        raw = self.excel([[146690], [datetime(2026, 6, 12)], ['VALID', '', '', 12345]])
        result = self.upload(raw)
        self.assertTrue(result.data['valid'])
        self.assertEqual(result.data['error_count'], 1)
        self.assertEqual(result.data['errors'][0]['row'], 3)
        self.assertIn('fecha', result.data['errors'][0]['message'])
        self.assertEqual(result.data['warning_count'], 2)
        self.assertIn('job_id', result.data)
        self.assertEqual(result.data['summary']['rows'], 2)
        self.assertEqual(result.data['pending_error_count'], 1)
        self.assertEqual(Part.objects.count(), 0)

    def test_confirmation_requires_same_file_user_and_unchanged_catalog(self):
        part = Part.objects.create(sku='EXISTING', name='BEFORE')
        raw = self.excel([['EXISTING', 'AFTER']])
        preview = self.upload(raw).data
        self.assertEqual(self.upload(raw, mode='import').status_code, 409)
        self.assertEqual(self.upload(self.excel([['OTHER']]), mode='import', preview_token=preview['preview_token']).status_code, 409)
        other_root = User.objects.create_superuser('other-root', 'other-root@example.invalid', 'OtherRoot938!')
        self.client.force_authenticate(other_root)
        self.assertEqual(self.upload(raw, mode='import', preview_token=preview['preview_token']).status_code, 409)
        self.client.force_authenticate(self.root)
        part.name = 'NEWER EDIT'
        part.save()
        self.assertEqual(self.upload(raw, mode='import', preview_token=preview['preview_token']).status_code, 409)
        part.refresh_from_db()
        self.assertEqual(part.name, 'NEWER EDIT')

    def test_failed_batch_rolls_back_skus_as_well_as_codes(self):
        raw = self.excel([['NEW', '', '', 'NEW-ALT']])
        preview = self.upload(raw).data
        with patch('mall.catalog_import.PartCode.objects.bulk_create', side_effect=IntegrityError('Concurrent duplicate')):
            self.assertEqual(self.upload(raw, mode='import', preview_token=preview['preview_token']).status_code, 409)
        self.assertEqual(Part.objects.count(), 0)
        self.assertEqual(PartCode.objects.count(), 0)

    def test_bad_excel_headers_and_empty_file_are_rejected(self):
        for raw in [b'not an excel file', self.excel([], headers=['SKU']), self.excel([['NEW']], headers=['UNKNOWN']), self.excel([['NEW', 'PRICE']], headers=['SKU', 'PRECIO'])]:
            self.assertEqual(self.upload(raw).status_code, 400)
        self.assertEqual(Part.objects.count(), 0)

    def test_batch_accepts_more_than_5000_rows_and_only_updates_supplied_fields(self):
        existing = Part.objects.create(sku='BATCH-0000', name='OLD', description='KEEP ME', active=False)
        raw = self.excel([[f'BATCH-{index:04}', 'new name' if index == 0 else '', '', f'ALT-{index:04}'] for index in range(6001)])
        preview = self.upload(raw)
        self.assertTrue(preview.data['valid'])
        self.assertEqual(preview.data['preview_count'], 6001)
        imported = self.upload(raw, mode='import', preview_token=preview.data['preview_token'])
        self.assertEqual(imported.status_code, 200)
        self.assertFalse(imported.data['imported'])
        self.assertEqual(Part.objects.count(), 500)
        url = f"{self.url}jobs/{preview.data['job_id']}/"
        for index in range(1, preview.data['progress']['total_batches']):
            imported = self.client.post(url, {'batch_index': index})
            self.assertEqual(imported.status_code, 200)
        self.assertTrue(imported.data['imported'])
        self.assertEqual(Part.objects.count(), 6001)
        self.assertEqual(PartCode.objects.count(), 6001)
        existing.refresh_from_db()
        self.assertEqual(existing.name, 'NEW NAME')
        self.assertEqual(existing.description, 'KEEP ME')
        self.assertFalse(existing.active)

    def test_durable_batches_are_atomic_owner_bound_and_retry_safe(self):
        raw = self.excel([[f'RESUME-{index:04}', '', '', f'CODE-{index:04}'] for index in range(1001)])
        preview = self.upload(raw)
        self.assertEqual(preview.data['progress']['total_batches'], 3)
        self.assertEqual(Part.objects.count(), 0)
        url = f"{self.url}jobs/{preview.data['job_id']}/"
        self.assertEqual(self.client.post(url, {'batch_index': 1}).status_code, 409)
        first = self.client.post(url, {'batch_index': 0})
        self.assertFalse(first.data['imported'])
        self.assertEqual(first.data['progress']['processed_skus'], 500)
        self.assertEqual(Part.objects.count(), 500)
        self.assertEqual(self.client.post(url, {'batch_index': 0}).data['progress']['completed_batches'], 1)
        self.assertEqual(Part.objects.count(), 500)
        with patch('mall.catalog_import.PartCode.objects.bulk_create', side_effect=IntegrityError('Concurrent duplicate')):
            self.assertEqual(self.client.post(url, {'batch_index': 1}).status_code, 409)
        self.assertEqual(Part.objects.count(), 500)
        self.assertEqual(PartCode.objects.count(), 500)
        self.assertEqual(self.client.get(url).data['progress']['completed_batches'], 1)
        other = User.objects.create_superuser('other-job-root', 'other-job-root@example.invalid', 'AnotherRoot938!')
        self.client.force_authenticate(other)
        self.assertEqual(self.client.get(url).status_code, 404)
        self.assertEqual(self.client.post(url, {'batch_index': 1}).status_code, 404)
        self.client.force_authenticate(self.user)
        self.assertEqual(self.client.get(url).status_code, 403)
        self.client.force_authenticate(self.root)
        self.assertEqual(self.client.post(url, {'batch_index': 1}).status_code, 200)
        completed = self.client.post(url, {'batch_index': 2})
        self.assertTrue(completed.data['imported'])
        self.assertEqual(completed.data['progress']['processed_skus'], 1001)
        self.assertEqual(completed.data['applied_summary']['created_skus'], 1001)
        self.assertEqual(Part.objects.count(), 1001)
        self.assertEqual(self.client.post(url, {'batch_index': 2}).status_code, 200)
        self.assertEqual(self.client.get(url).data['status'], 'completed')

    def test_shared_code_in_separate_brands_across_batches_does_not_invalidate_plan(self):
        rows = [[f'BRANDED-{index:04}'] for index in range(501)]
        rows[0] = ['BRANDED-0000', '', '', 'SHARED', 'OEM-A']
        rows[500] = ['BRANDED-0500', '', '', 'SHARED', 'OEM-B']
        preview = self.upload(self.excel(rows)).data
        self.assertTrue(preview['valid'])
        url = f"{self.url}jobs/{preview['job_id']}/"
        self.assertEqual(self.client.post(url, {'batch_index': 0}).status_code, 200)
        self.assertEqual(self.client.post(url, {'batch_index': 1}).status_code, 200)
        self.assertEqual(PartCode.objects.filter(code='SHARED').count(), 2)

    def test_categories_import_and_stale_later_batch_preserves_previous_commits(self):
        columns = self.columns + ['CATEGORIA', 'SUBCATEGORIA']
        rows = [[f'CATEG-{index:04}', '', '', '', '', '', 'frenos', 'discos'] for index in range(501)]
        preview = self.upload(self.excel(rows, columns)).data
        self.assertEqual(preview['preview'][0]['category'], 'FRENOS')
        url = f"{self.url}jobs/{preview['job_id']}/"
        self.assertEqual(self.client.post(url, {'batch_index': 0}).status_code, 200)
        part = Part.objects.get(sku='CATEG-0000')
        self.assertEqual(part.category, 'FRENOS')
        self.assertEqual(part.subcategory, 'DISCOS')
        Part.objects.create(sku='CATEG-0500')
        self.assertEqual(self.client.post(url, {'batch_index': 1}).status_code, 409)
        self.assertEqual(self.client.get(url).data['progress']['completed_batches'], 1)
        self.assertEqual(Part.objects.count(), 501)
        self.assertEqual(Part.objects.get(sku='CATEG-0500').category, '')

    def test_real_catalog_scale_can_be_previewed_without_a_row_cap(self):
        from .catalog_import import read_excel
        raw = self.excel([[f'SCALE-{index:05}'] for index in range(33320)])
        groups, errors, rows = read_excel(raw)
        self.assertEqual(rows, 33320)
        self.assertEqual(len(groups), 33320)
        self.assertEqual(errors, [])
