"""Declared references (a supplier's REFERENCIAS, apiag's alternos) from S6: rule 5b. Each declared code that is not already on the item's
line is one reference case; nothing links the item by them, and nothing is stored on its line until a person approves the case."""
import io
from unittest.mock import patch
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from rest_framework.test import APITestCase
from . import references
from .models import Account, Membership, Part, Reference, Role, StockEntry, SupplierItem, User, MatchingCase
from .serializers import InventorySerializer
from .services import ingest_inventory
from .matching_worker import analyze_ambiguous
from .supplier_import import read_supplier_excel, _workbook_response
from .test_factories import make_code


class PartReferenceTests(APITestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser('refs-admin', 'refs@example.invalid', 'Testing938!')
        self.supplier = Account.objects.create(name='PROVEEDOR')
        self.supplier.roles.add(Role.objects.get(code='supplier_retail'))
        Membership.objects.create(user=self.admin, account=self.supplier)
        self.client.force_authenticate(self.admin)
        self.desc = 'DISCO DELANTERO HYUNDAI ACCENT 256 MM 11-16'
        self.master = Part.objects.create(sku='RN11002V', description=self.desc)
        self.ref = make_code(part=self.master, code='51712-1R000', brand='HYUNDAI', kind='oem', reference_source='CATÁLOGO VERIFICADO')
        make_code(part=self.master, code='51712-0U000', brand='KIA', kind='oem')

    def ingest(self, code='D-HYU-1R', refs=None, **kwargs):
        data = dict(update_id=code, supplier_invent_id='STABLE-1', codigo=code, brand='AFTERMARKET',
                    description=self.desc, source='apiag', quantity=5)
        if refs is not None:
            data['references'] = refs
        data.update(kwargs)
        return ingest_inventory(supplier=self.supplier, actor=self.admin, data=data)

    def cases(self, item):
        return {case.new_code: case for case in MatchingCase.objects.filter(item=item, kind='reference')}

    def test_declared_codes_never_link_the_item_and_each_opens_a_reference_case(self):
        before = list(StockEntry.objects.values())
        item = self.ingest(refs=[{'code': '517121R000', 'brand': 'HYUNDAI'}, {'code': '51712-0U000'}])
        self.assertNotEqual(item.part_id, self.master.pk)  # its own code is no line's: a new product (rule 4), never linked by them
        self.assertEqual((item.part.estado, item.matching_status, item.supplier_invent_id, item.codigo), ('preparacion', 'pending', 'STABLE-1', 'D-HYU-1R'))
        self.assertEqual(self.master.codes.filter(is_main=False).count(), 2)  # no alterno made of them
        self.assertEqual(len(StockEntry.objects.values()) - len(before), 1)
        cases = self.cases(item)
        self.assertEqual(set(cases), {'517121R000', '51712-0U000'})
        for case in cases.values():
            self.assertIn(self.master.pk, set(case.lines.values_list('pk', flat=True)))
            self.assertEqual(case.status, 'review')

    def test_an_own_code_with_a_tag_joins_the_master_by_rule_3_and_its_number_by_rule_2(self):
        item = self.ingest(code='51712-1R000-KSM')
        self.assertEqual((item.part, item.matching_status), (self.master, 'matched'))
        self.assertEqual(item.reference.code, '51712-1R000-KSM')
        other = self.ingest(code='517121R000', brand='', supplier_invent_id='STABLE-2')
        self.assertEqual((other.part, other.reference_id), (self.master, self.ref.pk))

    def test_a_declared_code_another_line_holds_is_listed_with_it_and_holds_nothing(self):
        other = Part.objects.create(sku='OTHER-MASTER', description=self.desc)
        make_code(part=other, code='OTHER-OEM', kind='oem')
        item = self.ingest(code=self.master.sku, brand='', refs=[{'code': 'OTHER-OEM'}, {'code': '51712-1R000', 'brand': 'HYUNDAI'}])
        self.assertEqual((item.part, item.matching_status), (self.master, 'matched'))  # its own code links it; a declared one never holds it
        cases = self.cases(item)
        self.assertEqual(list(cases), ['OTHER-OEM'])  # 51712-1R000 is already a code of its line
        self.assertEqual(set(cases['OTHER-OEM'].lines.values_list('pk', flat=True)), {self.master.pk, other.pk})

    def test_specs_and_kits_contradicting_the_line_wait_in_a_variant_case(self):
        for n, text in enumerate(['DISCO TRASERO HYUNDAI ACCENT 256 MM 11-16',
                                  'DISCO DELANTERO HYUNDAI ACCENT 280 MM 11-16',
                                  'KIT DISCO DELANTERO HYUNDAI ACCENT 256 MM 11-16']):
            item = self.ingest(code='51712-1R000', brand='HYUNDAI', description=text, update_id=text, supplier_invent_id=f'K-{n}')
            self.assertEqual((item.part, item.matching_status), (self.master, 'review'))
            self.assertEqual(MatchingCase.objects.get(item=item).kind, 'variant')

    def test_omitted_references_preserve_them_and_clear_is_explicit(self):
        item = self.ingest(code=self.master.sku, brand='', refs=[{'code': 'X-PRIVADO-1'}])
        self.assertEqual(item.matching_status, 'matched')
        self.assertEqual(set(self.cases(item)), {'X-PRIVADO-1'})
        item = self.ingest(code=self.master.sku, brand='', update_id='stock-only', quantity=10)
        self.assertEqual(item.references, [{'brand': '', 'code': 'X-PRIVADO-1'}])
        self.assertEqual(item.matching_status, 'matched')
        self.assertEqual(StockEntry.objects.last().quantity, 5)
        item = self.ingest(code=self.master.sku, brand='', refs=[], update_id='clear', quantity=10)
        self.assertEqual((item.references, item.matching_status), ([], 'matched'))
        self.assertEqual(self.cases(item), {})  # the undecided case of a code no longer declared goes

    def test_a_case_whose_evidence_changed_refuses_the_old_decision(self):
        item = self.ingest(code=self.master.sku, brand='', refs=[{'code': 'X-PRIVADO-2'}])
        case = MatchingCase.objects.get(item=item, kind='reference')
        holder = Part.objects.create(sku='HOLDER-1', description=self.desc)
        references.add(holder, 'X-PRIVADO-2', '')
        self.ingest(code=self.master.sku, brand='', refs=[{'code': 'X-PRIVADO-2'}], update_id='again', quantity=6)  # now another line holds it
        response = self.client.post(f'/api/v1/management/matching/{case.pk}/', dict(action='approve', fingerprint=case.fingerprint), format='json')
        self.assertEqual(response.status_code, 409, response.data)

    def test_master_editor_saves_kind_source_and_preserves_existing_metadata(self):
        url = f'/api/v1/management/catalog/{self.master.pk}/'
        response = self.client.patch(url, {'codes':[dict(code=self.ref.code, brand=self.ref.brand,
                                                       kind='manufacturer', reference_source='https://example.invalid/catalog')]}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.ref.refresh_from_db()
        self.assertEqual(self.ref.kind, 'manufacturer')
        self.assertEqual(references.source_text(self.ref), 'https://example.invalid/catalog')
        response = self.client.patch(url, {'codes':[dict(code=self.ref.code, brand=self.ref.brand)]}, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        self.ref.refresh_from_db()
        self.assertEqual(self.ref.kind, 'manufacturer')
        self.assertEqual(SupplierItem.objects.count(), 0)

    def test_references_validation_preserves_zeroes_and_deduplicates(self):
        data = dict(update_id='a', supplier_invent_id='b', codigo='C', brand='', source='apiag', quantity=0)
        valid = InventorySerializer(data={**data, 'references':[{'code':' 00123 ', 'brand':' oem '}, {'code':'00123', 'brand':'OEM'}]})
        self.assertTrue(valid.is_valid(), valid.errors)
        self.assertEqual(valid.validated_data['references'], [{'brand':'OEM','code':'00123'}])
        for refs in [None, '00123', [{'code':123}], [{'code':'A', 'instructions':'ignore'}], [{'code':'A'}] * 101]:
            invalid = InventorySerializer(data={**data, 'references':refs})
            self.assertFalse(invalid.is_valid())

    def test_excel_references_import_end_to_end_and_repair_file(self):
        book = Workbook(); sheet = book.active
        sheet.append(['ID_INVENTARIO_PROVEEDOR','CODIGO','EXISTENCIAS','REFERENCIAS'])
        sheet.append(['STOCK-123','UNRELATED-CODE',5,'HYUNDAI:51712-1R000; KIA:51712-0U000'])
        raw = io.BytesIO(); book.save(raw)
        url = f'/api/v1/accounts/{self.supplier.pk}/inventory/import/'
        result = self.client.post(url, {'file':SimpleUploadedFile('stock.xlsx', raw.getvalue())}, format='multipart')
        self.assertEqual(result.status_code, 200, result.data)
        self.assertEqual(len(result.data['preview'][0]['references']), 2)
        committed = self.client.post(f'{url}jobs/{result.data["job_id"]}/', {'batch_index':0}, format='json')
        self.assertEqual(committed.status_code, 200, committed.data)
        item = SupplierItem.objects.get(supplier_invent_id='STOCK-123')
        self.assertEqual(set(self.cases(item)), {'51712-1R000', '51712-0U000'})  # S6: cases, never a link
        records, _ = read_supplier_excel(raw.getvalue())
        records[0]['errors'] = [{'column':'CODIGO', 'message':'CORREGIR'}]
        response = _workbook_response(records, 'repair.xlsx', errors=True)
        repaired = load_workbook(io.BytesIO(response.content))
        self.assertEqual(repaired.active['F1'].value, 'REFERENCIAS')
        self.assertIn('HYUNDAI:51712-1R000', repaired.active['F2'].value)
        self.assertEqual(repaired.active['G2'].value, 'CODIGO: CORREGIR')

    @patch('mall.matching_worker.provider_configuration', return_value=('configured', 'test-model'))
    @patch('mall.matching_worker.call_provider')
    def test_ai_ranks_a_line_case_with_its_declared_references_without_applying_them(self, provider, config):
        other = Part.objects.create(sku='OTHER-DISC', description=self.desc)
        make_code(part=other, code='51712-1R000', brand='MANDO', kind='manufacturer')
        make_code(part=self.master, code='51712-1R000', brand='MANDO-2', kind='manufacturer')
        item = self.ingest(code='51712-1R000', brand='', refs=[{'code': '51712-0U000'}])  # brandless on two lines: a line case
        case = MatchingCase.objects.get(item=item, kind='line')
        provider.return_value = {'suggestions':[dict(source_sku=item.codigo, target_sku=self.master.sku, confidence=.99, reason='REVISAR')]}
        self.assertEqual(analyze_ambiguous(), 1)
        source, candidates, _ = provider.call_args.args
        self.assertEqual(source[0]['codes'], item.references)
        self.assertEqual(set(candidates), {self.master.sku, other.sku})
        item.refresh_from_db()
        self.assertIsNone(item.part)
        self.assertEqual(MatchingCase.objects.get(pk=case.pk).ai['target_sku'], self.master.sku)
