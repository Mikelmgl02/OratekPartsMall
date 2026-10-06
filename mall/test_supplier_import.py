import io
from datetime import datetime, timedelta
from unittest.mock import patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.exceptions import ValidationError
from rest_framework.test import APITestCase

from .models import Account, InventoryUpdate, Membership, Part, PartCode, Role, StockEntry, SupplierItem, User
from .services import ingest_inventory
from .supplier_import import stock_plan
from .supplier_import_models import SupplierInventoryImportJob


class SupplierExcelImportTests(APITestCase):
    headers = ['ID_INVENTARIO_PROVEEDOR', 'CODIGO', 'MARCA', 'DESCRIPCION', 'EXISTENCIAS']

    def setUp(self):
        self.user = User.objects.create_user('supplier-import', 'supplier-import@example.invalid', 'Supplier938!')
        self.other = User.objects.create_user('supplier-other', 'supplier-other@example.invalid', 'Supplier938!')
        self.account = Account.objects.create(name='PROVEEDOR')
        self.account.roles.add(Role.objects.get(code='supplier_retail'))
        Membership.objects.create(user=self.user, account=self.account)
        Membership.objects.create(user=self.other, account=self.account)
        self.second_account = Account.objects.create(name='SEGUNDO PROVEEDOR')
        self.second_account.roles.add(Role.objects.get(code='supplier_wholesale'))
        Membership.objects.create(user=self.user, account=self.second_account)
        self.client.force_authenticate(self.user)
        self.part = Part.objects.create(sku='58411-1R000')
        PartCode.objects.create(part=self.part, code='58411-1R000-G')
        self.item = self.seed('STABLE-5', quantity=5)

    def seed(self, stable_id, *, quantity=5, source='upload', codigo='58411-1R000-G', brand='KSM', description='TAMBOR'):
        return ingest_inventory(supplier=self.account, actor=self.user,
                                data={'update_id': f'seed-{stable_id}', 'supplier_invent_id': stable_id,
                                      'codigo': codigo, 'brand': brand, 'description': description,
                                      'source': source, 'quantity': quantity})

    def base_url(self, account=None):
        return f'/api/v1/accounts/{(account or self.account).pk}/inventory/import/'

    def file(self, rows, *, headers=None, customize=None):
        book = Workbook()
        sheet = book.active
        sheet.title = 'EXISTENCIAS'
        sheet.append(headers or self.headers)
        for row in rows:
            sheet.append(row)
        if customize:
            customize(sheet)
        output = io.BytesIO()
        book.save(output)
        return SimpleUploadedFile('EXISTENCIAS.xlsx', output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def preview(self, rows, **kwargs):
        return self.client.post(self.base_url(), {'file': self.file(rows, **kwargs)}, format='multipart')

    def job_url(self, job_id, account=None):
        return f'{self.base_url(account)}jobs/{job_id}/'

    def commit(self, job_id, index=0, account=None):
        return self.client.post(self.job_url(job_id, account), {'batch_index': index}, format='json')

    def test_absolute_balances_create_positive_credit_and_debit_and_skip_unchanged_movements(self):
        zero = self.seed('ZERO-5', quantity=5)
        same = self.seed('SAME-5', quantity=5)
        omitted = self.seed('OMITTED', quantity=8)
        self.item.reserved_quantity = 2
        self.item.save()
        ledger_before = list(StockEntry.objects.values())
        result = self.preview([
            ['STABLE-5', '58411-1r000-g', 'ksm', 'tambor', 10],
            ['ZERO-5', '58411-1R000-G', 'KSM', 'TAMBOR', 0],
            ['SAME-5', '58411-1R000-G', 'KSM', 'TAMBOR', 5],
        ])
        self.assertEqual(result.status_code, 200)
        self.assertEqual(list(StockEntry.objects.values()), ledger_before)
        self.assertEqual(result.data['summary']['credit_units'], 5)
        self.assertEqual(result.data['summary']['debit_units'], 5)
        self.assertEqual([(row['direction'], row['movement_quantity'], row['new_quantity']) for row in result.data['preview']],
                         [('credit', 5, 10), ('debit', 5, 0), ('none', 0, 5)])
        saved = self.commit(result.data['job_id'])
        self.assertEqual(saved.status_code, 200)
        self.assertTrue(saved.data['imported'])
        self.assertEqual(saved.data['applied_summary']['unchanged_items'], 1)
        self.item.refresh_from_db()
        self.assertEqual((self.item.reported_quantity, self.item.reserved_quantity, self.item.part_id), (10, 2, self.part.pk))
        zero.refresh_from_db()
        same.refresh_from_db()
        omitted.refresh_from_db()
        self.assertEqual((zero.reported_quantity, same.reported_quantity, omitted.reported_quantity), (0, 5, 8))
        movements = list(StockEntry.objects.order_by('pk').values('quantity', 'direction', 'reported_after'))[-2:]
        self.assertEqual(movements, [{'quantity': 5, 'direction': 'credit', 'reported_after': 10},
                                     {'quantity': 5, 'direction': 'debit', 'reported_after': 0}])
        ledger_count = StockEntry.objects.count()
        again = self.commit(result.data['job_id'])
        self.assertEqual(again.status_code, 200)
        self.assertEqual(StockEntry.objects.count(), ledger_count)
        self.assertEqual(InventoryUpdate.objects.filter(key__startswith='excel-').count(), 3)

    def test_matching_uses_canonical_aliases_and_code_changes_require_review_without_changing_identity(self):
        different = Part.objects.create(sku='DIFFERENT')
        result = self.preview([
            ['STABLE-5', different.sku, 'KSM', 'NUEVO', 5],
            ['NEW-ALIAS', '58411-1R000-G', 'OTHER', 'NUEVO', 2],
            ['UNKNOWN', 'UNRECOGNIZED', '', 'SIN COINCIDENCIA', 0],
        ])
        self.assertEqual(result.status_code, 200)
        self.assertEqual([row['matching_status'] for row in result.data['preview']], ['review', 'matched', 'pending'])
        self.assertEqual(result.data['preview'][0]['part_sku'], self.part.sku)
        self.assertEqual(result.data['summary']['created_items'], 2)
        self.assertEqual(self.commit(result.data['job_id']).status_code, 200)
        self.item.refresh_from_db()
        self.assertEqual((self.item.part_id, self.item.matching_status, self.item.supplier_invent_id), (self.part.pk, 'review', 'STABLE-5'))
        self.assertEqual(SupplierItem.objects.get(supplier_invent_id='NEW-ALIAS').part_id, self.part.pk)
        self.assertIsNone(SupplierItem.objects.get(supplier_invent_id='UNKNOWN').part_id)
        self.assertEqual(Part.objects.count(), 2)

    def test_apiag_owned_rows_and_every_duplicate_stable_id_are_deferred(self):
        apiag = self.seed('ERP-OWNED', source='apiag', quantity=9)
        result = self.preview([
            ['ERP-OWNED', '58411-1R000-G', 'KSM', 'TAMBOR', 0],
            ['DUPLICATE', '58411-1R000-G', 'KSM', 'TAMBOR', 4],
            ['DUPLICATE', '58411-1R000-G', 'KSM', 'TAMBOR', 8],
            ['STABLE-5', '58411-1R000-G', 'KSM', 'TAMBOR', 6],
        ])
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.data['summary']['rejected_rows'], 3)
        self.assertEqual(result.data['summary']['valid_rows'], 1)
        self.assertEqual(result.data['error_count'], 3)
        self.assertEqual(self.commit(result.data['job_id']).status_code, 200)
        apiag.refresh_from_db()
        self.assertEqual((apiag.reported_quantity, apiag.source), (9, 'apiag'))
        self.assertFalse(SupplierItem.objects.filter(supplier_invent_id='DUPLICATE').exists())

    def test_quantity_errors_date_ids_and_original_formulas_survive_and_error_workbook_can_be_reuploaded(self):
        result = self.preview([
            ['BLANK', 'CODE', '', '', None],
            ['NEGATIVE', 'CODE', '', '', -1],
            ['FRACTION', 'CODE', '', '', 1.5],
            ['BOOLEAN', 'CODE', '', '', True],
            [datetime(2026, 6, 12), 'CODE', '', '', 2],
            ['FORMULA', 'CODE', '', '', '=1+1'],
            ['ERROR', 'CODE', '', '', '#VALUE!'],
        ])
        self.assertEqual(result.status_code, 200)
        self.assertFalse(result.data['valid'])
        self.assertEqual(result.data['status'], 'review')
        self.assertEqual(result.data['summary']['rejected_rows'], 7)
        job = SupplierInventoryImportJob.objects.get(pk=result.data['job_id'])
        self.assertEqual(job.rejected_rows[4]['original']['supplier_invent_id']['kind'], 'date')
        self.assertEqual(job.rejected_rows[5]['original']['quantity']['kind'], 'formula')
        job.expires_at = timezone.now() - timedelta(days=30)
        job.save()
        self.assertEqual(self.client.get(self.job_url(job.pk)).status_code, 200)
        download = self.client.get(f'{self.job_url(job.pk)}errors/')
        self.assertEqual(download.status_code, 200)
        workbook = load_workbook(io.BytesIO(download.content), data_only=False)
        sheet = workbook.active
        self.assertEqual(sheet['E7'].value, '=1+1')
        self.assertEqual(sheet['E7'].data_type, 's')
        sheet['A6'] = 'DATE-CORRECTED'
        for number in range(2, 9):
            sheet.cell(number, 5, 0)
        output = io.BytesIO()
        workbook.save(output)
        repair = self.client.post(self.base_url(), {'file': SimpleUploadedFile('CORREGIDO.xlsx', output.getvalue())}, format='multipart')
        self.assertTrue(repair.data['valid'])
        self.assertEqual(repair.data['summary']['valid_rows'], 7)
        self.assertEqual(self.commit(repair.data['job_id']).status_code, 200)
        self.assertEqual(SupplierItem.objects.filter(supplier=self.account).count(), 8)

    def test_safe_numeric_identifiers_and_leading_zeros_are_preserved(self):
        result = self.preview([[12, 34, '', '', 1], ['00Ab', '58411-1r000-g', '', '', 2]],
                              customize=lambda sheet: (setattr(sheet['A2'], 'number_format', '00000'), setattr(sheet['B2'], 'number_format', '0000')))
        self.assertEqual(result.status_code, 200)
        self.assertEqual([(row['supplier_invent_id'], row['codigo']) for row in result.data['preview']], [('00012', '0034'), ('00Ab', '58411-1R000-G')])
        self.assertEqual(self.commit(result.data['job_id']).status_code, 200)
        self.assertTrue(SupplierItem.objects.filter(supplier_invent_id='00012').exists())
        self.assertTrue(SupplierItem.objects.filter(supplier_invent_id='00Ab').exists())

    def test_absent_metadata_columns_preserve_existing_values_and_present_blanks_clear_them(self):
        result = self.preview([['STABLE-5', '58411-1R000-G', 8]], headers=['supplier_invent_id', 'CODE', 'STOCK'])
        self.assertEqual(result.status_code, 200)
        self.assertEqual((result.data['preview'][0]['brand'], result.data['preview'][0]['description']), ('KSM', 'TAMBOR'))
        self.assertEqual(self.commit(result.data['job_id']).status_code, 200)
        result = self.preview([['STABLE-5', '58411-1R000-G', '', '', 8]])
        self.assertEqual(self.commit(result.data['job_id']).status_code, 200)
        self.item.refresh_from_db()
        self.assertEqual((self.item.brand, self.item.description), ('', ''))

    def test_jobs_template_and_error_download_are_owner_and_supplier_scoped(self):
        job = self.preview([['STABLE-5', '58411-1R000-G', 'KSM', 'TAMBOR', 6]]).data['job_id']
        for method in ['get', 'post']:
            arguments = {'data': {'batch_index': 0}, 'format': 'json'} if method == 'post' else {}
            self.assertEqual(getattr(self.client, method)(self.job_url(job, self.second_account), **arguments).status_code, 404)
        self.client.force_authenticate(self.other)
        self.assertEqual(self.client.get(self.job_url(job)).status_code, 404)
        self.assertEqual(self.commit(job).status_code, 404)
        self.assertEqual(self.client.get(f'{self.job_url(job)}errors/').status_code, 404)
        self.client.force_authenticate(None)
        self.assertEqual(self.client.get(f'{self.base_url()}template/').status_code, 401)
        self.assertEqual(self.client.get(self.job_url(job)).status_code, 401)
        self.client.force_authenticate(self.user)
        self.account.roles.clear()
        self.assertEqual(self.client.get(f'{self.base_url()}template/').status_code, 403)
        self.assertEqual(self.commit(job).status_code, 403)
        self.item.refresh_from_db()
        self.assertEqual(self.item.reported_quantity, 5)

    def test_current_balance_source_reservations_and_mapping_changes_reject_stale_batches(self):
        mutations = [
            lambda: SupplierItem.objects.filter(pk=self.item.pk).update(reported_quantity=4),
            lambda: SupplierItem.objects.filter(pk=self.item.pk).update(reserved_quantity=1),
            lambda: SupplierItem.objects.filter(pk=self.item.pk).update(source='apiag'),
            lambda: SupplierItem.objects.filter(pk=self.item.pk).update(part=None, matching_status='pending'),
        ]
        for mutate in mutations:
            SupplierItem.objects.filter(pk=self.item.pk).update(reported_quantity=5, reserved_quantity=0, source='upload', part=self.part, matching_status='matched')
            result = self.preview([['STABLE-5', '58411-1R000-G', 'KSM', 'TAMBOR', 10]])
            mutate()
            count = StockEntry.objects.count()
            self.assertEqual(self.commit(result.data['job_id']).status_code, 409)
            self.assertEqual(StockEntry.objects.count(), count)
            self.assertEqual(SupplierInventoryImportJob.objects.get(pk=result.data['job_id']).completed_batches, 0)
        result = self.preview([['NEW-PART', 'NEW-CODE', '', '', 3]])
        PartCode.objects.create(part=self.part, code='NEW-CODE')
        self.assertEqual(self.commit(result.data['job_id']).status_code, 409)
        self.assertFalse(SupplierItem.objects.filter(supplier_invent_id='NEW-PART').exists())

    def test_stock_change_between_preview_plan_and_staging_is_detected_before_job_is_saved(self):
        calls = 0

        def concurrent_change(supplier, rows, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                SupplierItem.objects.filter(pk=self.item.pk).update(reported_quantity=7)
            return stock_plan(supplier, rows, **kwargs)

        with patch('mall.supplier_import.stock_plan', side_effect=concurrent_change):
            result = self.preview([['STABLE-5', '58411-1R000-G', 'KSM', 'TAMBOR', 3]])
        self.assertEqual(result.status_code, 409)
        self.assertEqual(SupplierInventoryImportJob.objects.count(), 0)
        self.assertEqual(StockEntry.objects.count(), 1)

    def test_one_failed_row_rolls_back_the_entire_batch_and_keeps_the_checkpoint(self):
        result = self.preview([['STABLE-5', '58411-1R000-G', 'KSM', 'TAMBOR', 10],
                               ['NEW', '58411-1R000-G', 'KSM', 'TAMBOR', 3]])
        calls = 0

        def fail_second(**kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise ValidationError('Conflicto simulado de existencias.')
            return ingest_inventory(**kwargs)

        with patch('mall.supplier_import.ingest_inventory', side_effect=fail_second):
            failed = self.commit(result.data['job_id'])
        self.assertEqual(failed.status_code, 400)
        self.item.refresh_from_db()
        self.assertEqual(self.item.reported_quantity, 5)
        self.assertFalse(SupplierItem.objects.filter(supplier_invent_id='NEW').exists())
        self.assertEqual(StockEntry.objects.count(), 1)
        self.assertEqual(SupplierInventoryImportJob.objects.get(pk=result.data['job_id']).completed_batches, 0)
        self.assertEqual(self.commit(result.data['job_id']).status_code, 200)

    def test_new_catalog_match_after_fingerprint_rolls_back_every_row_in_the_batch(self):
        result = self.preview([['STABLE-5', '58411-1R000-G', 'KSM', 'TAMBOR', 10],
                               ['PHANTOM', 'UNSEEN-CATALOG-CODE', '', 'PENDIENTE', 3]])
        self.assertEqual(result.data['preview'][1]['matching_status'], 'pending')

        def add_match_before_ingestion(**kwargs):
            if kwargs['data']['supplier_invent_id'] == 'PHANTOM':
                # Represents a concurrent creation of an absent matching SKU
                # after snapshot validation but before the matcher runs.
                Part.objects.create(sku='UNSEEN-CATALOG-CODE')
            return ingest_inventory(**kwargs)

        with patch('mall.supplier_import.ingest_inventory', side_effect=add_match_before_ingestion):
            response = self.commit(result.data['job_id'])
        self.assertEqual(response.status_code, 409)
        self.item.refresh_from_db()
        self.assertEqual(self.item.reported_quantity, 5)
        self.assertFalse(SupplierItem.objects.filter(supplier_invent_id='PHANTOM').exists())
        self.assertFalse(InventoryUpdate.objects.filter(key__startswith='excel-').exists())
        self.assertEqual(StockEntry.objects.count(), 1)
        job = SupplierInventoryImportJob.objects.get(pk=result.data['job_id'])
        self.assertEqual((job.completed_batches, job.processed_rows), (0, 0))
        self.assertFalse(Part.objects.filter(sku='UNSEEN-CATALOG-CODE').exists())

    def test_more_than_5000_rows_are_staged_and_checkpoint_retries_never_commit_the_next_batch(self):
        rows = [[f'LARGE-{index}', '58411-1R000-G', 'KSM', 'TAMBOR', 1] for index in range(5001)]
        result = self.preview(rows)
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.data['summary']['valid_rows'], 5001)
        self.assertEqual(result.data['progress']['total_batches'], 11)
        self.assertEqual(result.data['preview_count'], 5001)
        self.assertEqual(len(result.data['preview']), 100)
        self.assertEqual(SupplierItem.objects.count(), 1)
        first = self.commit(result.data['job_id'])
        self.assertEqual(first.status_code, 200)
        self.assertEqual(first.data['progress']['processed_rows'], 500)
        self.assertEqual(SupplierItem.objects.count(), 501)
        repeat = self.commit(result.data['job_id'])
        self.assertEqual(repeat.data['progress']['processed_rows'], 500)
        self.assertEqual(SupplierItem.objects.count(), 501)
        resume = self.client.get(self.job_url(result.data['job_id']))
        self.assertEqual(resume.data['progress']['next_batch'], 1)
        self.assertEqual(self.commit(result.data['job_id'], 2).status_code, 409)

    def test_expired_draft_cannot_commit_but_committed_retry_returns_saved_checkpoint(self):
        result = self.preview([['STABLE-5', '58411-1R000-G', 'KSM', 'TAMBOR', 6]])
        job = SupplierInventoryImportJob.objects.get(pk=result.data['job_id'])
        job.expires_at = timezone.now() - timedelta(days=1)
        job.save(update_fields=['expires_at'])
        self.assertEqual(self.commit(job.pk).status_code, 409)
        job.expires_at = timezone.now() + timedelta(days=1)
        job.save(update_fields=['expires_at'])
        self.assertEqual(self.commit(job.pk).status_code, 200)
        job.expires_at = timezone.now() - timedelta(days=1)
        job.save(update_fields=['expires_at'])
        self.assertEqual(self.commit(job.pk).status_code, 200)

    def test_template_uses_text_ids_and_bad_headers_and_files_are_global_errors(self):
        response = self.client.get(f'{self.base_url()}template/')
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.content))
        self.assertEqual(book.active['A2'].value, '000123')
        self.assertEqual(book.active['A2'].number_format, '@')
        for headers in [['ID', 'CODIGO'], ['ID', 'CODIGO', 'CANTIDAD', 'UNKNOWN'], ['ID', 'SUPPLIER_INVENT_ID', 'CODIGO', 'CANTIDAD']]:
            self.assertEqual(self.preview([['VALUE', 'CODE', 2]], headers=headers).status_code, 400)
        self.assertEqual(self.client.post(self.base_url(), {'file': SimpleUploadedFile('broken.xlsx', b'broken')}, format='multipart').status_code, 400)
        self.assertEqual(SupplierInventoryImportJob.objects.count(), 0)
