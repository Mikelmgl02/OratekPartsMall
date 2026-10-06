import io
import uuid
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from unittest import skipUnless

from django.db import IntegrityError, connection, transaction
from openpyxl import load_workbook
from rest_framework.test import APITestCase, APITransactionTestCase

from . import test_deals as deal_tests
from . import test_requests as request_tests
from .matching_models import MatchingQueue
from .models import InventoryUpdate, Membership, StockEntry, SupplierItem
from .pricing_models import PriceChange, PriceList, PriceListEntry, PricingAuditEvent, SupplierItemPricing, SupplierPricingSettings
from .supplier_import import stock_plan


class PriceListTests(APITestCase):
    setUp = request_tests.SupplierRequestWorkflowTests.setUp
    account = request_tests.SupplierRequestWorkflowTests.account
    item = request_tests.SupplierRequestWorkflowTests.item

    def base(self, account=None):
        return f'/api/v1/accounts/{(account or self.supplier_a).pk}'

    def as_user(self, user=None):
        self.client.force_authenticate(user or self.seller_a)
        return self.client

    def create(self, code='GENERAL', name='Lista general', user=None, **extra):
        return self.as_user(user).post(f'{self.base()}/price-lists/', {'code': code, 'name': name, **extra}, format='json')

    def update(self, price_list, version, user=None, **changes):
        return self.as_user(user).post(f'{self.base()}/price-lists/{price_list}/', {'expected_version': version, **changes}, format='json')

    def write(self, changes=(), items=(), user=None, account=None):
        body = {**({'changes': list(changes)} if changes else {}), **({'items': list(items)} if items else {})}
        return self.as_user(user).post(f'{self.base(account)}/prices/', body, format='json')

    def change(self, price_list, item, price, revision=None):
        return {'list_id': str(price_list), 'supplier_item_id': str(item.pk), 'unit_price': price, 'expected_revision': revision}

    def lists(self):
        self.create()
        mayorista = self.create('MAYORISTA', 'Mayoristas').data['id']
        return PriceList.objects.get(code='GENERAL').pk, mayorista

    def test_lists_are_private_versioned_and_keep_exactly_one_active_default(self):
        empty = self.as_user().get(f'{self.base()}/price-lists/')
        self.assertEqual(empty.data, {'can_configure': True, 'item_count': 2, 'results': []})
        first = self.create(' general ', ' lista general ')
        self.assertEqual((first.status_code, first.data['code'], first.data['name'], first.data['currency'], first.data['is_default'], first.data['version']),
                         (201, 'GENERAL', 'LISTA GENERAL', 'USD', True, 1))
        self.assertEqual(self.create(' general ', ' lista general ').data, first.data)
        self.assertEqual(PriceList.objects.count(), 1)
        reused = self.create('GENERAL', 'OTRO NOMBRE')
        self.assertEqual((reused.status_code, str(reused.data['code'][0])), (400, 'Ya tienes una lista con el código GENERAL.'))
        for code in ['MAYOR ISTA', 'NIVEL-A', 'Ñ', 'X' * 31, '', 'MINIMO', ' piso ']:
            with self.subTest(code=code):
                self.assertEqual(self.create(code, 'LISTA').status_code, 400)
        SupplierPricingSettings.objects.create(supplier=self.supplier_a, default_currency='PAB')
        nivel = self.create('nivel_a', 'Nivel A', is_default=True)
        self.assertEqual((nivel.data['code'], nivel.data['currency'], nivel.data['is_default']), ('NIVEL_A', 'PAB', True))
        self.assertEqual(list(PriceList.objects.filter(is_default=True).values_list('code', flat=True)), ['NIVEL_A'])
        self.assertEqual(str(self.create('MINIMO', 'X').data['code'][0]), 'El código MINIMO está reservado para el precio mínimo (PRECIO_MINIMO); elige otro.')
        # The demoted default moves its version too: an edit made from the old values cannot silently take the default back.
        general = PriceList.objects.get(code='GENERAL')
        self.assertEqual((general.version, general.is_default), (2, False))
        reclaimed = self.update(general.pk, 1, name='LISTA GENERAL', is_default=True, active=True, currency='USD')
        self.assertEqual((reclaimed.status_code, reclaimed.data['price_list']['is_default']), (409, False))
        renamed = self.update(general.pk, 2, name='lista base')
        self.assertEqual((renamed.status_code, renamed.data['name'], renamed.data['version']), (200, 'LISTA BASE', 3))
        stale = self.update(general.pk, 2, name='otra')
        self.assertEqual((stale.status_code, stale.data['detail'], stale.data['price_list']['name']), (409, 'Otro usuario actualizó esta lista. Revisa sus valores actuales.', 'LISTA BASE'))
        self.assertEqual(self.update(general.pk, 2, name='LISTA BASE').data['version'], 3, 'Equal values are a no-op even with an old version.')
        nivel_id = nivel.data['id']
        for changes, field in [({'active': False}, 'active'), ({'is_default': False}, 'is_default')]:
            with self.subTest(changes=changes):
                refused = self.update(nivel_id, 1, **changes)
                self.assertEqual((refused.status_code, list(refused.data)), (400, [field]))
        self.assertEqual(self.update(general.pk, 3, is_default=True, active=False).status_code, 400)
        self.assertEqual(self.update(general.pk, 3, active=False).data['active'], False)
        self.assertEqual(self.update(general.pk, 4, is_default=True).status_code, 400, 'An archived list cannot become the default.')
        moved = self.update(general.pk, 4, is_default=True, active=True)
        self.assertEqual((moved.data['is_default'], moved.data['active']), (True, True))
        self.assertEqual(self.update(nivel_id, 1, active=False).status_code, 409, 'NIVEL_A moved to version 2 when it stopped being the default.')
        self.assertEqual(self.update(nivel_id, 2, active=False).data['active'], False)
        self.assertEqual(self.update(general.pk, 5, currency='PAB').data['currency'], 'PAB')
        self.assertEqual(self.write([self.change(general.pk, self.item_a, '10.00')]).status_code, 200)
        locked = self.update(general.pk, 7, currency='USD')
        self.assertEqual((locked.status_code, str(locked.data['currency'][0])), (400, 'No puedes cambiar la moneda de una lista que ya tiene precios.'))
        listed = {row['code']: row for row in self.as_user().get(f'{self.base()}/price-lists/').data['results']}
        self.assertEqual([(code, row['priced_count'], row['missing_count']) for code, row in listed.items()], [('GENERAL', 1, 1), ('NIVEL_A', 0, 2)])
        self.assertEqual(PricingAuditEvent.objects.filter(kind='price_list_changed').count(), 7)
        # The database backs the service: never two defaults, never an archived default.
        with self.assertRaises(IntegrityError), transaction.atomic():
            PriceList.objects.create(supplier=self.supplier_a, code='OTRA', name='OTRA', is_default=True, created_by=self.seller_a)
        with self.assertRaises(IntegrityError), transaction.atomic():
            PriceList.objects.filter(code='GENERAL').update(active=False)

    def test_list_endpoints_follow_membership_and_configuration_permission(self):
        general, _ = self.lists()
        self.assertEqual(self.as_user(self.seller_b).get(f'{self.base()}/price-lists/').status_code, 404)
        self.assertEqual(self.as_user(self.root).get(f'{self.base()}/price-lists/').status_code, 404)
        self.assertEqual(self.as_user(self.buyer).get(f'{self.base(self.client_account)}/price-lists/').status_code, 403)
        self.assertEqual(self.as_user(self.seller_b).post(f'{self.base(self.supplier_b)}/price-lists/{general}/', {'expected_version': 1, 'name': 'X'},
                                                          format='json').status_code, 404)
        SupplierPricingSettings.objects.create(supplier=self.supplier_a, config_min_permission='manager')
        listed = self.as_user().get(f'{self.base()}/price-lists/')
        self.assertEqual((listed.status_code, listed.data['can_configure']), (200, False))
        for response in [self.create('NUEVA', 'NUEVA'), self.update(general, 1, name='X'), self.write([self.change(general, self.item_a, '1.00')])]:
            self.assertEqual(response.status_code, 403)
        Membership.objects.filter(user=self.seller_a).update(permission='manager')
        self.assertEqual(self.create('NUEVA', 'NUEVA').status_code, 201)

    def test_bulk_edits_are_compare_and_set_all_or_nothing_and_retry_safe(self):
        general, mayorista = self.lists()
        first = self.write([self.change(general, self.item_a, '15.00'), self.change(mayorista, self.item_a, '13.50'), self.change(general, self.item_a2, '8')],
                           [{'supplier_item_id': str(self.item_a.pk), 'discount_group': ' filtros ', 'floor_price': '12.00', 'expected_revision': None}])
        self.assertEqual(first.status_code, 200, first.data)
        self.assertEqual(first.data['summary'], {'created': 3, 'updated': 0, 'removed': 0, 'unchanged': 0, 'floor_updates': 1, 'line_updates': 1})
        row = next(row for row in first.data['rows'] if row['supplier_item_id'] == str(self.item_a.pk))
        self.assertEqual((row['discount_group'], row['floor_price'], row['item_pricing_revision'], set(row['prices'])), ('FILTROS', '12.00', 1, {'GENERAL', 'MAYORISTA'}))
        self.assertEqual((row['prices']['GENERAL']['unit_price'], row['prices']['GENERAL']['revision']), ('15.00', 1))
        self.assertEqual(PriceChange.objects.count(), 5)
        self.assertEqual(dict(PriceList.objects.values_list('code', 'version')), {'GENERAL': 2, 'MAYORISTA': 2})
        event = PricingAuditEvent.objects.get(kind='prices_edited')
        self.assertEqual((event.payload['created'], event.payload['lists'], event.payload['reference'][:3]), (3, {'GENERAL': 2, 'MAYORISTA': 1}, 'op:'))
        # A retry after a lost response (same values, old expectations) changes nothing and writes no history.
        retry = self.write([self.change(general, self.item_a, '15.00'), self.change(mayorista, self.item_a, '13.50'), self.change(general, self.item_a2, '8.00')],
                           [{'supplier_item_id': str(self.item_a.pk), 'discount_group': 'FILTROS', 'floor_price': '12', 'expected_revision': None}])
        self.assertEqual(retry.data['summary'], {'created': 0, 'updated': 0, 'removed': 0, 'unchanged': 4, 'floor_updates': 0, 'line_updates': 0})
        self.assertEqual((PriceChange.objects.count(), PricingAuditEvent.objects.filter(kind='prices_edited').count()), (5, 1))
        self.assertEqual(dict(PriceList.objects.values_list('code', 'version')), {'GENERAL': 2, 'MAYORISTA': 2})
        updated = self.write([self.change(general, self.item_a, '16.25', 1)])
        self.assertEqual((updated.data['summary']['updated'], updated.data['rows'][0]['prices']['GENERAL']['revision']), (1, 2))
        change = PriceChange.objects.filter(kind='list_price').first()
        self.assertEqual((change.old_value, change.new_value, change.source, change.actor_id), (Decimal('15.00'), Decimal('16.25'), 'manual', self.seller_a.pk))
        # A stale revision refuses the whole request, including its valid rows.
        conflict = self.write([self.change(general, self.item_a, '17.00', 1), self.change(general, self.item_a2, '9.00', 1)],
                              [{'supplier_item_id': str(self.item_a2.pk), 'floor_price': '1.00', 'expected_revision': 3}])
        self.assertEqual(conflict.status_code, 409)
        self.assertEqual(conflict.data['detail'], 'Otro usuario cambió algunos precios.')
        self.assertEqual(conflict.data['conflicts'], [{'supplier_item_id': str(self.item_a.pk), 'list_id': str(general), 'current': {'unit_price': '16.25', 'revision': 2}},
                                                      {'supplier_item_id': str(self.item_a2.pk), 'list_id': None,
                                                       'current': {'discount_group': '', 'floor_price': None, 'revision': None}}])
        self.assertEqual({row['supplier_item_id'] for row in conflict.data['rows']}, {str(self.item_a.pk), str(self.item_a2.pk)})
        self.assertEqual(PriceListEntry.objects.get(price_list=general, item=self.item_a2).unit_price, Decimal('8.00'))
        self.assertEqual(PriceChange.objects.count(), 6)
        # null removes the entry; a later re-creation continues the revision count, so the old revision can never match again.
        removed = self.write([self.change(general, self.item_a, None, 2)])
        self.assertEqual((removed.data['summary']['removed'], removed.data['rows'][0]['prices'].get('GENERAL')), (1, None))
        self.assertFalse(PriceListEntry.objects.filter(price_list=general, item=self.item_a).exists())
        self.assertEqual(self.write([self.change(general, self.item_a, None, 2)]).data['summary']['unchanged'], 1)
        self.assertEqual(self.write([self.change(general, self.item_a, '5.00', 2)]).status_code, 409)
        recreated = self.write([self.change(general, self.item_a, '5.00')])
        self.assertEqual(recreated.data['rows'][0]['prices']['GENERAL']['revision'], 4)
        cleared = self.write(items=[{'supplier_item_id': str(self.item_a.pk), 'discount_group': '', 'floor_price': None, 'expected_revision': 1}])
        self.assertEqual((cleared.data['summary']['line_updates'], cleared.data['summary']['floor_updates'], cleared.data['rows'][0]['item_pricing_revision']), (1, 1, 2))
        self.assertEqual(list(PriceChange.objects.filter(kind='discount_group').values_list('old_text', 'new_text')), [('FILTROS', ''), ('', 'FILTROS')])

    def test_price_values_are_validated_strictly_and_scoped_to_the_supplier(self):
        general, _ = self.lists()
        cases = [('12.345', 'Usa un máximo de dos decimales.'), ('-1.00', 'El precio de lista debe ser mayor que 0.'), ('0', 'El precio de lista debe ser mayor que 0.'),
                 ('12345678901', 'Usa como máximo 10 dígitos enteros.'), (True, 'Escribe un importe válido, por ejemplo 12.50.'), ('12,50', 'Escribe un importe válido, por ejemplo 12.50.')]
        for value, message in cases:
            with self.subTest(value=value):
                response = self.write([self.change(general, self.item_a, value)])
                self.assertEqual((response.status_code, str(response.data['changes'][0]['unit_price'][0])), (400, message))
        floor = self.write(items=[{'supplier_item_id': str(self.item_a.pk), 'floor_price': '-0.01', 'expected_revision': None}])
        self.assertEqual(str(floor.data['items'][0]['floor_price'][0]), 'El precio mínimo no puede ser negativo.')
        self.assertEqual(self.write(items=[{'supplier_item_id': str(self.item_a.pk), 'floor_price': '0', 'expected_revision': None}]).status_code, 200)
        foreign = self.write([self.change(general, self.item_b, '1.00')])
        self.assertEqual((foreign.status_code, foreign.data), (400, ['El artículo no pertenece a tu inventario.']))
        foreign = self.write(items=[{'supplier_item_id': str(self.item_b.pk), 'discount_group': 'X', 'expected_revision': None}])
        self.assertEqual(foreign.data, ['El artículo no pertenece a tu inventario.'])
        other_list = PriceList.objects.create(supplier=self.supplier_b, code='AJENA', name='AJENA', is_default=True, created_by=self.seller_b)
        self.assertEqual(self.write([self.change(other_list.pk, self.item_a, '1.00')]).data, ['Una de las listas de precios no existe.'])
        for body in [{}, {'changes': []}, {'changes': [self.change(general, self.item_a, '1.00')] * 2},
                     {'items': [{'supplier_item_id': str(self.item_a.pk), 'expected_revision': None}]},
                     {'changes': [self.change(general, self.item_a, '1.00', True)]},
                     {'changes': [self.change(general, self.item_a, '1.00')] * 501}]:
            with self.subTest(body=str(body)[:80]):
                self.assertEqual(self.as_user().post(f'{self.base()}/prices/', body, format='json').status_code, 400)
        self.assertEqual(PriceListEntry.objects.count(), 0)
        self.assertEqual(self.as_user(self.buyer).post(f'{self.base(self.client_account)}/prices/', {}, format='json').status_code, 403)
        self.assertEqual(self.as_user(self.seller_b).post(f'{self.base()}/prices/', {}, format='json').status_code, 404)

    def test_price_grid_filters_search_and_item_history(self):
        general, mayorista = self.lists()
        self.write([self.change(general, self.item_a, '9.00'), self.change(mayorista, self.item_a2, '7.00')],
                   [{'supplier_item_id': str(self.item_a.pk), 'discount_group': 'TAMBORES', 'floor_price': '10.00', 'expected_revision': None}])
        page = self.as_user().get(f'{self.base()}/prices/')
        self.assertEqual((page.data['count'], [row['supplier_invent_id'] for row in page.data['results']]), (2, ['A-001', 'A-002']))
        self.assertEqual(set(page.data['results'][0]), {'supplier_item_id', 'supplier_invent_id', 'codigo', 'brand', 'description', 'matching_status',
                                                        'discount_group', 'floor_price', 'item_pricing_revision', 'prices', 'updated_at'})
        def ids(query):
            return [row['supplier_invent_id'] for row in self.as_user().get(f'{self.base()}/prices/?{query}').data['results']]
        self.assertEqual(ids('status=priced'), ['A-001'])
        self.assertEqual(ids('status=missing'), ['A-002'])
        self.assertEqual(ids('status=below_floor'), ['A-001'])
        self.assertEqual(ids(f'status=priced&list={mayorista}'), ['A-002'])
        self.assertEqual(ids(f'status=below_floor&list={mayorista}'), [])
        self.assertEqual(ids('search=tambores'), ['A-001'])
        self.assertEqual(ids('search=a-002'), ['A-002'])
        self.assertEqual(self.as_user().get(f'{self.base()}/prices/?status=caro').status_code, 400)
        self.assertEqual(self.as_user().get(f'{self.base()}/prices/?list={uuid.uuid4()}').status_code, 404)
        self.write([self.change(general, self.item_a, '9.50', 1)])
        history = self.as_user().get(f'{self.base()}/prices/{self.item_a.pk}/history/')
        self.assertEqual(history.data['count'], 4)
        latest = history.data['results'][0]
        self.assertEqual({key: latest[key] for key in ['kind', 'price_list', 'old_value', 'new_value', 'source', 'actor']},
                         {'kind': 'list_price', 'price_list': {'id': str(general), 'code': 'GENERAL'}, 'old_value': '9.00', 'new_value': '9.50',
                          'source': 'manual', 'actor': {'name': 'request-seller-a'}})
        self.assertEqual(self.as_user().get(f'{self.base()}/prices/{self.item_b.pk}/history/').status_code, 404)
        self.assertEqual(self.as_user(self.seller_b).get(f'{self.base(self.supplier_b)}/prices/{self.item_a.pk}/history/').status_code, 404)
        audit = self.as_user().get(f'{self.base()}/pricing/history/')
        self.assertEqual([row['kind'] for row in audit.data['results']], ['prices_edited', 'prices_edited', 'price_list_changed', 'price_list_changed'])
        self.assertEqual(self.as_user().get(f'{self.base()}/pricing/history/?search=MAYORISTA').data['count'], 0)
        self.assertEqual(self.as_user().get(f'{self.base()}/pricing/history/?search=price_list').data['count'], 2)
        self.assertEqual(self.as_user(self.seller_b).get(f'{self.base()}/pricing/history/').status_code, 404)

    def test_price_writes_never_touch_stock_rows_catalog_freshness_or_staged_stock_imports(self):
        general, _ = self.lists()
        rows = [{'row': 2, 'data': {'supplier_invent_id': 'A-001', 'codigo': '58411-1R000-G', 'brand': 'KSM', 'quantity': 4}}]
        plan, fingerprint = stock_plan(self.supplier_a, [dict(row, data=dict(row['data'])) for row in rows])
        self.assertFalse(plan['errors'])
        before = (list(SupplierItem.objects.order_by('pk').values()), StockEntry.objects.count(), InventoryUpdate.objects.count(), list(MatchingQueue.objects.values()))
        self.client.force_authenticate(self.buyer)
        catalog = self.client.get('/api/v1/catalog/').data['results'][0]['availability']
        self.assertEqual(self.write([self.change(general, self.item_a, '9.00')],
                                    [{'supplier_item_id': str(self.item_a.pk), 'discount_group': 'X', 'floor_price': '1.00', 'expected_revision': None}]).status_code, 200)
        self.assertEqual(self.write([self.change(general, self.item_a, None, 1)]).status_code, 200)
        self.assertEqual((list(SupplierItem.objects.order_by('pk').values()), StockEntry.objects.count(), InventoryUpdate.objects.count(), list(MatchingQueue.objects.values())), before)
        self.assertEqual(stock_plan(self.supplier_a, [dict(row, data=dict(row['data'])) for row in rows])[1], fingerprint)
        self.client.force_authenticate(self.buyer)
        self.assertEqual(self.client.get('/api/v1/catalog/').data['results'][0]['availability'], catalog)
        self.assertEqual(SupplierItemPricing.objects.get().supplier_id, self.supplier_a.pk)

    def test_export_round_trips_the_import_columns_privately(self):
        general, mayorista = self.lists()
        archived = self.create('VIEJA', 'VIEJA').data['id']
        self.update(archived, 1, active=False)
        self.write([self.change(general, self.item_a, '15.00'), self.change(mayorista, self.item_a, '13.5'), self.change(general, self.item_a2, '1234567.89')],
                   [{'supplier_item_id': str(self.item_a.pk), 'discount_group': 'FILTROS', 'floor_price': '12.00', 'expected_revision': None}])
        SupplierItem.objects.filter(pk=self.item_a2.pk).update(supplier_invent_id='000123')
        response = self.as_user().get(f'{self.base()}/prices/export/')
        self.assertEqual((response.status_code, response['Content-Type'], response['Cache-Control']),
                         (200, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'private, no-store'))
        sheet = load_workbook(io.BytesIO(response.content))['PRECIOS']
        rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
        self.assertEqual(rows[0], ['ID_INVENTARIO_PROVEEDOR', 'CODIGO', 'MARCA', 'DESCRIPCION', 'LINEA', 'PRECIO_MINIMO', 'PRECIO', 'PRECIO_MAYORISTA'])
        self.assertEqual(rows[1:], [['000123', '58411-1R000-KSM', 'KSM', 'TAMBOR A-002', None, None, 1234567.89, None],
                                    ['A-001', '58411-1R000-G', 'KSM', 'TAMBOR A-001', 'FILTROS', 12, 15, 13.5]])
        only = load_workbook(io.BytesIO(self.as_user().get(f'{self.base()}/prices/export/?list={mayorista}').content))['PRECIOS']
        self.assertEqual([cell.value for cell in next(only.iter_rows())][-1], 'PRECIO_MAYORISTA')
        self.assertEqual(self.as_user().get(f'{self.base()}/prices/export/?list={uuid.uuid4()}').status_code, 400)
        self.assertEqual(self.as_user(self.buyer).get(f'{self.base(self.client_account)}/prices/export/').status_code, 403)
        self.assertEqual(self.as_user(self.seller_b).get(f'{self.base()}/prices/export/').status_code, 404)
        self.assertEqual(self.as_user(self.root).get(f'{self.base()}/prices/export/').status_code, 404)


@skipUnless(connection.vendor == 'postgresql', 'Row locks are only meaningful on PostgreSQL.')
class PriceWriteConcurrencyTests(APITransactionTestCase):
    # Transaction tests flush migrated roles, so they reuse the self-contained deal fixture.
    setUp = deal_tests.ConcurrentDealTests.setUp
    call = deal_tests.ConcurrentDealTests.call

    def test_parallel_edits_of_the_same_revision_give_one_winner(self):
        general = PriceList.objects.create(supplier=self.supplier, code='GENERAL', name='GENERAL', is_default=True, created_by=self.seller)
        url = f'/api/v1/accounts/{self.supplier.pk}/prices/'
        def edit(price, revision=1):
            return self.call(self.seller, url, {'changes': [{'list_id': str(general.pk), 'supplier_item_id': str(self.item.pk), 'unit_price': price,
                                                             'expected_revision': revision}]})[0]
        self.assertEqual(edit('10.00', None), 200)
        with ThreadPoolExecutor(max_workers=2) as pool:
            self.assertEqual(sorted(pool.map(edit, ['11.00', '12.00'])), [200, 409])
        self.assertEqual(PriceListEntry.objects.get().revision, 2)
        self.assertEqual(PriceChange.objects.count(), 2)
